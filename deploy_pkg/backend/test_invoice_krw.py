"""거래명세서(1실) 원화 표기 회귀 테스트 — 두 백엔드 모두.

화면에서는 행별 RATE 칸이 비어 있으면 상단 '환율' 값을 쓴다
(MicronInvoice.js: `const itemRate = Number(it.rate) || rate`).
엑셀·PDF 생성도 같아야 한다 — 안 그러면 화면엔 원화가 보이는데
받은 파일에는 외화만 찍힌다.

main.py(로컬 인메모리) / main_aws.py(EB 배포) 가 갈라지지 않게 둘 다 건다.
"""
import io

import openpyxl
import pytest
from pypdf import PdfReader

import main
import main_aws

# 프론트가 실제로 보내는 모양: 행 rate 는 빈 문자열, 환율은 최상위 rate 하나
PAYLOAD = {
    "items": [{"part": "PIC16F1708-I/ML", "qty": "1000", "price": "1.5", "rate": ""}],
    "customer": "METACOM",
    "date": "2026-09-07",
    "rate": 1400,
    "issue_date": "2026-09-07",
    "person_in_charge": "test@unitrontech.com",
}

BACKENDS = [pytest.param(main, id="main"), pytest.param(main_aws, id="main_aws")]


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_has_krw_columns(mod):
    ws = openpyxl.load_workbook(io.BytesIO(mod._build_invoice_xlsx_bytes(PAYLOAD))).active
    rate, price_krw, amount_krw = ws["F14"].value, ws["G14"].value, ws["H14"].value
    assert rate == 1400, f"RATE 칸이 비었다: {rate!r}"
    assert price_krw == 2100, f"U/PRICE(₩) 가 비었다: {price_krw!r}"
    assert amount_krw == 2100000, f"AMOUNT(₩) 가 비었다: {amount_krw!r}"


@pytest.mark.parametrize("mod", BACKENDS)
def test_pdf_has_krw_amounts(mod):
    pdf = mod._build_invoice_pdf_bytes(PAYLOAD)
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "2,100,000" in text, "PDF 품목행에 원화 금액이 없다"
    assert "2,310,000" in text, "PDF 합계에 원화 총액(VAT 포함)이 없다"


def test_row_rate_still_wins():
    """행별 RATE 를 직접 넣은 경우는 그 값이 우선한다."""
    payload = dict(PAYLOAD, items=[dict(PAYLOAD["items"][0], rate="1300")])
    ws = openpyxl.load_workbook(io.BytesIO(main_aws._build_invoice_xlsx_bytes(payload))).active
    assert ws["F14"].value == 1300
    assert ws["H14"].value == 1950000
