"""발주요청서 보고 양식(표1·표2·표3) 생성 — main.py(로컬) · main_aws.py(배포) 공용.

거래명세서와 같은 이유로 한 곳에 모았다: 두 백엔드가 사본을 따로 들면 반드시 갈라진다.

양식은 사용자가 직접 손본 `발주요청서_보고_2026-09-08.xlsx` 의 빨간 코멘트를 기준으로 한다.
색·글자크기·병합·숫자서식을 그 파일에서 그대로 옮겼으므로 임의로 바꾸지 말 것.
"""
import io
import math
from datetime import date, datetime

import pandas as pd

PO_REPORT_T1_COLS = [
    "CPO", "Order Date", "Unitrontech CRD", "고객사 요청일",
    "담당 Sales", "End Customer", "DID", "MPN", "Package",
    "Qty", "DCPL", "Resale Price", "Resale AMT", "PO Customer", "Remark",
]

# 표1 헤더: 1·2행 병합이 기본이고 'Unitrontech / CRD' 만 2단으로 쌓는다
T1_HEADER_TOP = [
    "CPO", "Order Date", "Unitrontech", "고객사 요청일", "담당 Sales", "End Customer",
    "DID", "MPN", "Package", "Qty", "DCPL", "Resale Price", "Resale AMT",
    "PO Customer", "Remark",
]
T1_STACKED_COL = 3          # C열만 1행 Unitrontech / 2행 CRD
T1_GREEN_COLS = (3, 11, 14)  # C(Unitrontech CRD) · K(DCPL) · N(PO Customer)

T2_COLS = ["Date", "Sales", "PART NO.", "Customer", "기 수주", "신규 수주",
           "Inventory", "Backlog",
           "매출 M", "매출 +1M", "매출 +2M", "매출 +3M", "매출 ~+4M", "Remarks"]
T2_SUB = {5: "수량", 6: "수량", 7: "Total", 8: "Total"}  # E~H 는 2행에 소제목

# 원본 파일에서 그대로 옮긴 서식
FMT_QTY = "#,##0"
FMT_USD = r'\$#,##0.00_);[Red]\(\$#,##0.00\)'
FMT_ACCT = r'_-* #,##0_-;\-* #,##0_-;_-* "-"_-;_-@_-'
FMT_DATE = "yyyy-mm-dd"     # 값은 데이터 AE열 그대로, 표기는 같은 표의 CRD 와 맞춘다

# ARGB 8자리로 둔다 — 6자리로 주면 openpyxl 이 알파를 00 으로 써서 원본과 코드가 달라진다
C_YELLOW, C_GREEN, C_BLUE, C_LIME, C_MONTH, C_BAL = (
    "FFFFE699", "FFC6E0B4", "FF95B3D7", "FF92D050", "FFBDD7EE", "FFFFFF00")
C_RED = "FFFF0000"   # 표3 New PO — 자동으로 채워진 값임을 표시

MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

T3_ROWS = ["Delivery", "Inventory", "Backlog", "New PO", "Balance"]
T3_AUTO_ROW = "New PO"   # 이 행만 자동 — 나머지는 사용자가 손으로 채운다
T3_TAIL_MONTHS = 6       # New PO 가 끝나는 달 + 6개월까지 보여준다

# 표4 (2026-09-10 추가요청): 월별 발주 수량.
# "대량 발주 건들은 CRD를 같은달에 쪼개 발주 넣는 경우가 있어 월별 발주 수량 파악을 위해"
# → 값은 표3 New PO 와 같은 수치(표1 Qty 를 CRD 월로 합산)를 품목×고객 행으로 눕힌 것.
T4_COLS = ["Part (MPN)", "End Customer", "DID"]
T4_TAIL_MONTHS = 6       # 첫 발주 달 + 6개월 (발주가 더 뒤까지 있으면 그 달까지)
C_T4_HDR = "FF99CCFF"    # 요청서 표4 헤더 채움색

# 회계연도 = 마이크론 기준 9월 시작. FQ1=Sep~Nov, FQ2=Dec~Feb, FQ3=Mar~May, FQ4=Jun~Aug.
# 연도는 종료연도 표기 — Sep 2026 은 YR2027. (요청서 표3 FQ 행에서 역산:
#  Mar~May26=FQ3 YR2026, Jun~Aug26=FQ4 YR2026, Sep~Nov26=FQ1 YR2027, Dec26~Feb27=FQ2 YR2027)
FISCAL_START_MONTH = 9


# ---------- 값 정리 ----------

def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.strftime("%Y-%m-%d")
    return v


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(v)
        return None if math.isnan(f) or math.isinf(f) else f
    except (ValueError, TypeError):
        return None


def _to_date(v):
    """문자열/타임스탬프 → date. 못 읽으면 None."""
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return pd.Timestamp(v).date()
    except Exception:
        return None


def _ym(d):
    """date → 연*12+월 정수. 월 계산을 연도 넘어가도 안전하게 하려고 쓴다."""
    return d.year * 12 + (d.month - 1)


def month_abbr(ym):
    return MONTH_ABBR[ym % 12]


def sales_plan_labels(base):
    """표2 '매출 계획' 5칸 라벨. Date 기준으로 월 약어를 자동 기입한다.

    2026-09-08 → M(Sep) / +1M(Oct) / +2M(Nov) / +3M(Dec) / ~+4M (Jan~)
    """
    if base is None:
        return ["M", "+1M", "+2M", "+3M", "~+4M"]
    b = _ym(base)
    return [f"M({month_abbr(b)})", f"+1M({month_abbr(b + 1)})",
            f"+2M({month_abbr(b + 2)})", f"+3M({month_abbr(b + 3)})",
            f"~+4M ({month_abbr(b + 4)}~)"]


def month_window(t1_rows):
    """표3 월 컬럼 범위 = Order date 달 ~ (New PO 마지막 달 + 6개월).

    New PO 는 표1 CRD 월에 꽂히므로, 마지막 CRD 달이 곧 New PO 마지막 달이다.
    """
    starts = [_to_date(r.get("Order Date")) for r in t1_rows]
    starts = [d for d in starts if d]
    crds = [_to_date(r.get("Unitrontech CRD")) for r in t1_rows]
    crds = [d for d in crds if d]
    if not starts and not crds:
        return []
    first = min(_ym(d) for d in (starts or crds))
    last_po = max(_ym(d) for d in (crds or starts))
    end = last_po + T3_TAIL_MONTHS
    if end < first:
        end = first
    return list(range(first, end + 1))


# ---------- 표1 / 표2 ----------

def build_preview(contents: bytes) -> dict:
    """발주요청서 데이터 양식 → 보고 양식(표1·표2) 데이터."""
    try:
        df = pd.read_excel(io.BytesIO(contents), header=0)
    except Exception:
        return {"error": "엑셀을 읽을 수 없습니다."}

    t1_rows = []
    for _, row in df.iterrows():
        if row.isna().all():
            continue
        rec = {
            "CPO": _clean(row.get("CUST PO#")),
            "Order Date": _clean(row.get("PO Date")),
            "Unitrontech CRD": _clean(row.get("CRD")),
            # 고객사 요청일: 데이터 양식 AE열(Customer SRD) 날짜를 그대로 쓴다.
            # 예전엔 '2027년 2월' 로 뭉갰는데 사용자가 원본 날짜를 요청했다.
            "고객사 요청일": _clean(row.get("Customer SRD")),
            "담당 Sales": _clean(row.get("FSE")),
            "End Customer": _clean(row.get("End customer")),
            "DID": _clean(row.get("DID")),
            "MPN": _clean(row.get("MPN")),
            "Package": _clean(row.get("BOX_TYPE")),
            "Qty": _clean(row.get("QTY")),
            "DCPL": _clean(row.get("DCPL")),
            "Resale Price": _clean(row.get("SP ($)")),
            "Resale AMT": _clean(row.get("Sales Amt ($)")),
            "PO Customer": _clean(row.get("PO Customer")),
            "Remark": None,  # 수동 입력 칸
        }
        if not rec.get("MPN") and not rec.get("DID"):
            continue
        if rec["Resale AMT"] in (None, 0, ""):
            qty, price = _num(rec["Qty"]), _num(rec["Resale Price"])
            if qty is not None and price is not None:
                rec["Resale AMT"] = round(qty * price, 2)
        t1_rows.append(rec)

    sum_qty = sum(_num(r.get("Qty")) or 0 for r in t1_rows)
    sum_amt = sum(_num(r.get("Resale AMT")) or 0 for r in t1_rows)

    t2_groups, t2_order = {}, []
    for r in t1_rows:
        key = (r.get("MPN"), r.get("End Customer"))
        if key not in t2_groups:
            t2_groups[key] = {
                "Date": r.get("Order Date"), "Sales": r.get("담당 Sales"),
                "PART NO.": r.get("MPN"), "Customer": r.get("End Customer"),
                "기 수주": None, "신규 수주": 0, "Inventory": None, "Backlog": None,
                "매출 M": None, "매출 +1M": None, "매출 +2M": None,
                "매출 +3M": None, "매출 ~+4M": None, "Remarks": None,
            }
            t2_order.append(key)
        t2_groups[key]["신규 수주"] += _num(r.get("Qty")) or 0

    return {
        "t1_columns": PO_REPORT_T1_COLS,
        "t1_rows": t1_rows,
        "t1_sum": {"Qty": round(sum_qty), "Resale AMT": round(sum_amt, 2)},
        "t2_rows": [t2_groups[k] for k in t2_order],
    }


# ---------- 표3 ----------

def build_t3_blocks(t1_rows):
    """표1에서 (DID)MPN 블록과 New PO 수량을 뽑는다.

    New PO 는 표1 CRD 의 '월' 버킷에 수량을 더한다. 나머지 행은 사용자가 채운다.
    """
    months = month_window(t1_rows)
    blocks, order = {}, []
    for r in t1_rows:
        did, mpn = r.get("DID"), r.get("MPN")
        if not mpn and not did:
            continue
        key = (did, mpn)
        if key not in blocks:
            title = f"({did}){mpn}" if did else str(mpn)
            blocks[key] = {"title": title, "new_po": {}}
            order.append(key)
        crd = _to_date(r.get("Unitrontech CRD"))
        qty = _num(r.get("Qty"))
        if crd and qty:
            ym = _ym(crd)
            blocks[key]["new_po"][ym] = blocks[key]["new_po"].get(ym, 0) + qty
    return months, [blocks[k] for k in order]


def fiscal_label(ym):
    """연*12+월 → "FQ3 YR2026". 월이 어디로 오든 라벨이 자동으로 나온다."""
    y, m = ym // 12, ym % 12 + 1
    idx = (m - FISCAL_START_MONTH) % 12          # Sep=0 … Aug=11
    fy = y + 1 if m >= FISCAL_START_MONTH else y  # 종료연도 표기
    return f"FQ{idx // 3 + 1} YR{fy}"


def fiscal_spans(months):
    """같은 FQ 끼리 묶는다 → [(시작 인덱스, 칸 수, 라벨)]. 표3 FQ 행 병합에 쓴다."""
    spans = []
    for i, ym in enumerate(months):
        lab = fiscal_label(ym)
        if spans and spans[-1][2] == lab:
            spans[-1][1] += 1
        else:
            spans.append([i, 1, lab])
    return [tuple(x) for x in spans]


# ---------- 표4 ----------

def t4_month_window(t1_rows):
    """표4 월 범위 = 첫 발주(CRD) 달 ~ 첫 달+6개월. 발주가 더 뒤까지 있으면 마지막 발주 달까지.

    요청서 예시(발주 Jun~Sep 2026)가 Jun~Dec 7칸이라 그 규칙을 그대로 쓴다.
    뒤를 자르지 않으므로 발주가 7개월을 넘겨도 안 잘린다.
    """
    crds = [_to_date(r.get("Unitrontech CRD")) for r in t1_rows]
    crds = [_ym(d) for d in crds if d]
    if not crds:
        return []
    first, last = min(crds), max(crds)
    return list(range(first, max(first + T4_TAIL_MONTHS, last) + 1))


def build_t4_rows(t1_rows):
    """(MPN, End Customer, DID) 별 월간 발주 수량. 같은 달에 쪼개 들어온 건은 한 칸으로 합친다."""
    months = t4_month_window(t1_rows)
    rows, order = {}, []
    for r in t1_rows:
        mpn, cust, did = r.get("MPN"), r.get("End Customer"), r.get("DID")
        if not mpn and not did:
            continue
        key = (mpn, cust, did)
        if key not in rows:
            rows[key] = {"Part (MPN)": mpn, "End Customer": cust, "DID": did, "qty": {}}
            order.append(key)
        crd = _to_date(r.get("Unitrontech CRD"))
        qty = _num(r.get("Qty"))
        if crd and qty:
            ym = _ym(crd)
            rows[key]["qty"][ym] = rows[key]["qty"].get(ym, 0) + qty
    return months, [rows[k] for k in order]


# ---------- 엑셀 ----------

def build_xlsx_bytes(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    t1_rows = data.get("t1_rows", [])
    t1_sum = data.get("t1_sum", {})
    t2_rows = data.get("t2_rows", [])

    def fill(rgb):
        return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")

    # 표1~4 테두리 (2026-09-10 요청). 요청서에서 직접 그려 보낸 선이 thin 이라 그대로 맞춘다.
    _side = Side(style="thin")
    THIN = Border(left=_side, right=_side, top=_side, bottom=_side)

    def box(ws, r1, c1, r2, c2):
        """직사각형 범위에 테두리. 병합칸도 구성 셀마다 넣어야 선이 끊기지 않는다."""
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = THIN

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wb = Workbook()

    # ===== 표1 =====
    ws1 = wb.active
    ws1.title = "표1"
    f_hdr1 = Font(name="맑은 고딕", size=10, bold=True)
    for j, h in enumerate(T1_HEADER_TOP, start=1):
        bg = fill(C_GREEN if j in T1_GREEN_COLS else C_YELLOW)
        top = ws1.cell(row=1, column=j, value=h)
        top.font, top.alignment, top.fill = f_hdr1, center, bg
        bottom = ws1.cell(row=2, column=j)
        bottom.font, bottom.alignment, bottom.fill = f_hdr1, center, bg
        if j == T1_STACKED_COL:
            bottom.value = "CRD"
        else:
            ws1.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
    ws1.row_dimensions[2].height = 17.25

    body = Font(name="맑은 고딕", size=10)
    for i, r in enumerate(t1_rows, start=3):
        for j, col in enumerate(PO_REPORT_T1_COLS, start=1):
            v = r.get(col)
            if col == "고객사 요청일":
                v = _to_date(v) or v
            c = ws1.cell(row=i, column=j, value=v)
            c.font = body
            if col == "고객사 요청일":
                c.number_format = FMT_DATE
            elif col == "Qty":
                c.number_format = FMT_QTY
            elif col in ("DCPL", "Resale Price", "Resale AMT"):
                c.number_format = FMT_USD

    sum_row = len(t1_rows) + 3
    sum_fill = fill("FFFFF2CC")
    for col_idx, key, fmt in ((10, "Qty", FMT_QTY), (13, "Resale AMT", FMT_USD)):
        c = ws1.cell(row=sum_row, column=col_idx, value=t1_sum.get(key))
        c.fill, c.font, c.number_format = sum_fill, Font(name="맑은 고딕", size=10, bold=True), fmt

    box(ws1, 1, 1, sum_row, len(T1_HEADER_TOP))

    for letter, w in zip("ABCDEFGHIJKLMNO",
                         [14, 12, 14, 14, 12, 18, 8, 28, 12, 10, 12, 12, 12, 18, 7.25]):
        ws1.column_dimensions[letter].width = w

    # ===== 표2 =====
    ws2 = wb.create_sheet("표2")
    f_hdr2 = Font(name="굴림", size=10, bold=True)
    base = _to_date(t2_rows[0].get("Date")) if t2_rows else None
    labels = sales_plan_labels(base)
    top_row = (["Date", "Sales", "PART NO.", "Customer", "기 수주", "신규 수주",
                "Inventory", "Backlog"] + ["매출 계획"] + [None] * 4 + ["Remarks"])
    for j, h in enumerate(top_row, start=1):
        bg = fill(C_LIME if 9 <= j <= 13 or j == 14 else C_BLUE)
        c = ws2.cell(row=1, column=j, value=h)
        c.font, c.alignment, c.fill = f_hdr2, center, bg
    ws2.merge_cells(start_row=1, start_column=9, end_row=1, end_column=13)  # 매출 계획
    for j in range(1, 15):
        if j in T2_SUB:
            c = ws2.cell(row=2, column=j, value=T2_SUB[j])
            c.font, c.alignment, c.fill = f_hdr2, center, fill(C_BLUE)
        elif 9 <= j <= 13:
            c = ws2.cell(row=2, column=j, value=labels[j - 9])
            c.font, c.alignment, c.fill = f_hdr2, center, fill(C_LIME)
        else:
            ws2.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
    ws2.row_dimensions[1].height = 17.25
    ws2.row_dimensions[2].height = 24.75

    # 본문 글꼴은 맑은 고딕 10pt (2026-09-10 요청). 헤더는 요청서대로 굴림을 유지한다.
    body2 = Font(name="맑은 고딕", size=10)
    # 빈칸에도 회계 서식을 미리 입혀 나중에 숫자를 넣어도 콤마가 붙게 한다
    for i, r in enumerate(t2_rows, start=3):
        for j, col in enumerate(T2_COLS, start=1):
            v = r.get(col)
            if col == "기 수주":
                # =Inventory Total + Backlog Total. G·H 를 손으로 채우면 즉시 계산된다.
                v = f"=G{i}+H{i}"
            c = ws2.cell(row=i, column=j, value=v)
            c.font = body2
            if 5 <= j <= 13:
                c.number_format = FMT_ACCT
    box(ws2, 1, 1, max(2, len(t2_rows) + 2), len(T2_COLS))

    for letter, w in zip("ABCDEFGHIJKLMN",
                         [12, 10, 28, 18, 10, 10, 10, 10, 10, 10, 10, 10, 12, 16]):
        ws2.column_dimensions[letter].width = w

    # ===== 표3 =====
    ws3 = wb.create_sheet("표3")
    months, blocks = build_t3_blocks(t1_rows)
    f_title = Font(name="맑은 고딕", size=9, bold=True)
    f_plan = Font(name="맑은 고딕", size=10, bold=True)
    body3 = Font(name="맑은 고딕", size=10)
    ncol = len(months)

    f_fq = Font(name="맑은 고딕", size=10)
    spans = fiscal_spans(months)

    row = 1
    for blk in blocks:
        # 제목 / FQ / 월 3단 머리 (FQ 행은 2026-09-10 요청으로 추가)
        head, fqrow, mrow = row, row + 1, row + 2
        t = ws3.cell(row=head, column=2, value=blk["title"])
        t.font, t.alignment = f_title, Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws3.merge_cells(start_row=head, start_column=2, end_row=mrow, end_column=2)
        if ncol:
            p = ws3.cell(row=head, column=3, value="소요계획")
            p.font, p.alignment = f_plan, center
            ws3.merge_cells(start_row=head, start_column=3, end_row=head, end_column=2 + ncol)
        for start, n, label in spans:
            c = ws3.cell(row=fqrow, column=3 + start, value=label)
            c.font, c.alignment = f_fq, center
            if n > 1:
                ws3.merge_cells(start_row=fqrow, start_column=3 + start,
                                end_row=fqrow, end_column=3 + start + n - 1)
        for k, ym in enumerate(months):
            c = ws3.cell(row=mrow, column=3 + k, value=month_abbr(ym))
            c.font, c.alignment, c.fill = body3, center, fill(C_MONTH)

        first_data = mrow + 1
        red3 = Font(name="맑은 고딕", size=10, color=C_RED)
        for n, label in enumerate(T3_ROWS):
            r = first_data + n
            is_auto = label == T3_AUTO_ROW
            lab = ws3.cell(row=r, column=2, value=label)
            lab.font = red3 if is_auto else body3
            lab.alignment = center   # 2026-09-10 요청: 항목 글씨 가운데 정렬
            is_balance = label == "Balance"
            if is_balance:
                lab.fill = fill(C_BAL)
            for k, ym in enumerate(months):
                col = 3 + k
                letter = get_column_letter(col)
                if is_balance:
                    # =Inventory-Delivery+Backlog+New PO — 손으로 채우면 즉시 계산된다
                    v = (f"={letter}{first_data + 1}-{letter}{first_data}"
                         f"+{letter}{first_data + 2}+{letter}{first_data + 3}")
                elif label == T3_AUTO_ROW:
                    v = blk["new_po"].get(ym)
                else:
                    v = None  # 수동 입력
                c = ws3.cell(row=r, column=col, value=v)
                c.font = red3 if is_auto else body3
                c.alignment, c.number_format = center, FMT_ACCT
                if is_balance:
                    c.fill = fill(C_BAL)
        box(ws3, head, 2, first_data + len(T3_ROWS) - 1, 2 + ncol)
        row = first_data + len(T3_ROWS) + 1  # 블록 사이 한 줄 비움

    ws3.column_dimensions["A"].width = 23.75
    ws3.column_dimensions["B"].width = 24
    for k in range(ncol):
        ws3.column_dimensions[get_column_letter(3 + k)].width = 10

    # ===== 표4 =====
    # 월별 발주 수량. 서식은 요청서 "추가요청 표4" 시트를 그대로 옮겼다.
    ws4 = wb.create_sheet("표4")
    t4_months, t4_rows = build_t4_rows(t1_rows)
    f_hdr4 = Font(name="맑은 고딕", size=10, bold=True)
    body4 = Font(name="맑은 고딕", size=10)
    hdr4 = T4_COLS + [month_abbr(ym) for ym in t4_months]
    for j, h in enumerate(hdr4, start=1):
        c = ws4.cell(row=1, column=j, value=h)
        c.font, c.alignment, c.fill = f_hdr4, center, fill(C_T4_HDR)
    for i, r in enumerate(t4_rows, start=2):
        for j, col in enumerate(T4_COLS, start=1):
            c = ws4.cell(row=i, column=j, value=r.get(col))
            c.font = body4
        for k, ym in enumerate(t4_months):
            # 발주 없는 달은 0 — 회계 서식이 "-" 로 보여준다 (요청서와 같은 모양)
            c = ws4.cell(row=i, column=len(T4_COLS) + 1 + k, value=r["qty"].get(ym, 0))
            c.font, c.number_format = body4, FMT_ACCT
    if t4_rows:
        box(ws4, 1, 1, len(t4_rows) + 1, len(hdr4))
    for letter, w in zip("ABC", [27.9, 17.4, 5.9]):
        ws4.column_dimensions[letter].width = w
    for k in range(len(t4_months)):
        ws4.column_dimensions[get_column_letter(len(T4_COLS) + 1 + k)].width = 8.8

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
