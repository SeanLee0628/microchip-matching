"""거래명세서 업로드 → 엑셀/PDF 회귀 테스트.

핵심 약속 세 가지를 건다.
1. 파일에 적힌 값을 서버가 다시 계산하지 않는다.
2. 업체별 소수점 표기(셀 서식)를 그대로 물려준다 — 2자리/4자리가 섞여 들어온다.
3. 외화만 / 원화만 쓰는 건은 반대쪽 통화 칸을 비운다.

main.py(로컬) / main_aws.py(배포) 가 또 갈라지지 않게 둘 다 건다.
"""
import io
import zipfile
from urllib.parse import unquote

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader

import main
import main_aws
from invoice_upload import decimals_from_fmt, parse_invoice_upload

BACKENDS = [pytest.param(main, id="main"), pytest.param(main_aws, id="main_aws")]

FMT_USD = r"\$#,##0.00"
FMT_RATE = r'_-* #,##0.00_-;\-* #,##0.00_-;_-* "-"_-;_-@_-'
FMT_KRW2 = r'"₩"#,##0.00_);[Red]\("₩"#,##0.00\)'
FMT_KRW4 = r'"₩"#,##0.0000_);[Red]\("₩"#,##0.0000\)'
FMT_KRW0 = r'"₩"#,##0_);[Red]\("₩"#,##0\)'

HEADERS = ["DATE", "SALES", "CUSTOMER", "MPN", "Q'ty", "U/PRICE($)", "AMOUNT($)",
           "RATE", "U/PRICE(\\)", "AMOUNT(\\)", "TAX(\\)", "Total AMT(\\)",
           "문서번호", "거래명세서 일자"]

# 실제 업로드 파일과 같은 모양 — 2자리 건, 원화 전용 건, 4자리 2품목 건
ROWS = [
    # 양쪽 표기 · 원화단가 소수점 2자리
    ["2026-08-26", "SW", "DS Tech", "MTFC256GBCAQTC-IT", 2000, 181.54, 363080,
     1383.1, 251088, 502176000, 50217600, 552393600, "2026-00322", "2026-08-26"],
    # 원화만 (RATE·U/PRICE($) 없음)
    ["2026-08-03", "Eric", "LGVS", "MT29GZ6A6BPIET-046AAT.112", 500, 0, 0,
     None, 28830, 14415000, 1441500, 15856500, "2026-00221", "2026-08-03"],
    # 같은 문서번호 2줄 · 원화단가 소수점 4자리
    ["2026-08-19", "KATE", "RS Automation", "MT28EW128ABA1HJS-0SIT", 1152, 19, 21888,
     1497.43, 28451.17, 32775747.84, 3277575, 36053322.84, "2026-00295", "2026-08-21"],
    ["2026-08-19", "KATE", "RS Automation", "MT41K64M16TW-107 AIT:J", 2000, 15.7, 31400,
     1497.43, 23509.651, 47019302, 4701930, 51721232, "2026-00295", "2026-08-21"],
]

# 열별 서식: 3~4행의 원화 단가만 4자리
COL_FMT = {6: FMT_USD, 7: FMT_USD, 8: FMT_RATE, 9: FMT_KRW2,
           10: FMT_KRW0, 11: FMT_KRW0, 12: FMT_KRW0}


def _sample_bytes(as_formula=False):
    """업로드 데이터 파일을 만든다. as_formula=True 면 캐시값 없는 수식만 넣는다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for i, r in enumerate(ROWS):
        row = list(r)
        if as_formula:
            n = i + 2
            row[6] = f"=E{n}*F{n}"          # AMOUNT($)
            row[9] = f"=E{n}*I{n}"          # AMOUNT(₩)
            row[10] = f"=ROUND(J{n}*0.1,0)"  # TAX(₩)
            row[11] = f"=J{n}+K{n}"         # Total AMT(₩)
        ws.append(row)
    for n in range(2, len(ROWS) + 2):
        for col, fmt in COL_FMT.items():
            ws.cell(row=n, column=col).number_format = fmt
        if n >= 4:  # RS Automation 두 줄은 원화 단가 4자리
            ws.cell(row=n, column=9).number_format = FMT_KRW4
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _invoices(mod, contents=None):
    res = TestClient(mod.app).post(
        "/api/invoice/upload-preview",
        files={"file": ("up.xlsx", contents or _sample_bytes(), "application/xlsx")},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert "error" not in body, body
    return body["invoices"]


def _sheet(mod, inv):
    return openpyxl.load_workbook(io.BytesIO(mod._build_invoice_xlsx_bytes(inv))).active


def _pdf_text(mod, inv):
    return PdfReader(io.BytesIO(mod._build_invoice_pdf_bytes(inv))).pages[0].extract_text()


# ---------- 파싱 ----------

def test_groups_by_doc_no():
    invs = parse_invoice_upload(_sample_bytes())["invoices"]
    assert [i["doc_no"] for i in invs] == ["2026-00322", "2026-00221", "2026-00295"]
    assert [i["item_count"] for i in invs] == [1, 1, 2], "같은 문서번호는 한 장이어야 한다"


def test_krw_only_row_detected():
    lgvs = parse_invoice_upload(_sample_bytes())["invoices"][1]
    assert lgvs["items"][0]["currency"] == "KRW"
    assert lgvs["totals"]["sub_usd"] is None, "원화 전용 건은 $ 합계가 비어야 한다"


def test_issue_date_and_person_come_from_file():
    rs = parse_invoice_upload(_sample_bytes())["invoices"][2]
    assert rs["issue_date"] == "2026-08-21", "거래명세서 일자"
    assert rs["date"] == "2026-08-19", "출고일자는 가장 이른 날"
    assert rs["person_in_charge"] == "KATE"


def test_formula_only_file_recovers_values():
    """Excel 로 저장된 적 없는 파일은 캐시값이 없다 — 수식대로 되살려야 한다."""
    invs = parse_invoice_upload(_sample_bytes(as_formula=True))["invoices"]
    it = invs[0]["items"][0]
    assert it["amount_usd"] == pytest.approx(363080), it
    assert it["amount_krw"] == pytest.approx(502176000), it
    assert it["tax_krw"] == pytest.approx(50217600), it


@pytest.mark.parametrize("fmt,expected", [
    (FMT_KRW2, 2), (FMT_KRW4, 4), (FMT_KRW0, 0), (FMT_USD, 2), ("General", None),
])
def test_decimals_from_fmt(fmt, expected):
    assert decimals_from_fmt(fmt) == expected


# ---------- 엑셀 ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_keeps_source_values(mod):
    """환산·반올림을 다시 하지 않는다 — 파일 값이 그대로 찍혀야 한다."""
    ws = _sheet(mod, _invoices(mod)[2])
    assert ws["G14"].value == pytest.approx(28451.17)
    assert ws["G15"].value == pytest.approx(23509.651)
    assert ws["H15"].value == pytest.approx(47019302)
    assert ws["F14"].value == pytest.approx(1497.43)


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_keeps_source_number_format(mod):
    """업체별 소수점 표기(2자리 / 4자리)를 셀 서식 그대로 물려준다."""
    assert _sheet(mod, _invoices(mod)[0])["G14"].number_format == FMT_KRW2
    ws = _sheet(mod, _invoices(mod)[2])
    assert ws["G14"].number_format == FMT_KRW4
    assert ws["G15"].number_format == FMT_KRW4


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_krw_only_leaves_usd_blank(mod):
    ws = _sheet(mod, _invoices(mod)[1])
    assert ws["G14"].value == pytest.approx(28830), "원화 단가는 있어야 한다"
    for cell in ("D14", "E14", "F14"):
        assert ws[cell].value is None, f"{cell} 에 외화가 찍혔다: {ws[cell].value!r}"
    for row in (24, 25, 26):  # 소계 / 부가세 / 합계
        assert ws.cell(row=row, column=5).value is None, f"{row}행 $ 합계가 비어야 한다"


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_totals_come_from_file(mod):
    ws = _sheet(mod, _invoices(mod)[2])
    assert ws.cell(row=24, column=8).value == pytest.approx(79795049.84)
    assert ws.cell(row=25, column=8).value == pytest.approx(7979505)
    assert ws.cell(row=26, column=8).value == pytest.approx(87774554.84)


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_has_doc_no(mod):
    ws = _sheet(mod, _invoices(mod)[0])
    assert ws["K4"].value == "문서번호"
    assert ws["L4"].value == "2026-00322"
    assert ws["L5"].value == "2026-08-26", "발행일은 문서번호 바로 아래"


@pytest.mark.parametrize("mod", BACKENDS)
def test_xlsx_bigo_stays_empty(mod):
    """적용환율 기준이 건마다 달라 비고에 고정 문구를 찍지 않는다."""
    ws = _sheet(mod, _invoices(mod)[0])
    assert not ws["A27"].value, f"비고가 채워졌다: {ws['A27'].value!r}"


# ---------- PDF ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_pdf_keeps_source_decimals(mod):
    t = _pdf_text(mod, _invoices(mod)[2])
    assert "₩28,451.1700" in t, "4자리 표기가 유지돼야 한다"
    assert "₩23,509.6510" in t
    assert "₩32,775,748" in t, "금액은 원본 서식(0자리)대로 반올림"


@pytest.mark.parametrize("mod", BACKENDS)
def test_pdf_krw_only_hides_usd(mod):
    t = _pdf_text(mod, _invoices(mod)[1])
    assert "₩28,830.00" in t
    assert "공급가액 합계($)" not in t, "원화 전용 건에 $ 합계가 나오면 안 된다"
    assert "공급가액 합계(₩)" in t


@pytest.mark.parametrize("mod", BACKENDS)
def test_pdf_has_doc_no_and_no_bigo(mod):
    t = _pdf_text(mod, _invoices(mod)[0])
    assert "문서번호 : 2026-00322" in t
    assert "최초매매기준율" not in t and "최초고시" not in t


# ---------- ZIP ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_zip_has_xlsx_and_pdf_per_invoice(mod):
    invs = _invoices(mod)
    res = TestClient(mod.app).post("/api/invoice/upload-zip", json={"invoices": invs})
    assert res.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(res.content)).namelist()
    assert len(names) == len(invs) * 2, names
    assert "거래명세서_RS Automation_2026-00295.xlsx" in names
    assert "거래명세서_RS Automation_2026-00295.pdf" in names


@pytest.mark.parametrize("mod", BACKENDS)
def test_zip_filename_header_is_latin1_safe(mod):
    res = TestClient(mod.app).post("/api/invoice/upload-zip",
                                   json={"invoices": _invoices(mod)})
    cd = res.headers["content-disposition"]
    cd.encode("latin-1")  # 한글을 그대로 넣으면 500 이 난다
    assert "거래명세서" in unquote(cd.split("filename*=UTF-8''")[1])


@pytest.mark.parametrize("mod", BACKENDS)
def test_zip_rejects_empty(mod):
    res = TestClient(mod.app).post("/api/invoice/upload-zip", json={"invoices": []})
    assert res.status_code == 400


@pytest.mark.parametrize("mod", BACKENDS)
def test_upload_preview_rejects_junk(mod):
    res = TestClient(mod.app).post(
        "/api/invoice/upload-preview",
        files={"file": ("x.xlsx", b"not an excel file", "application/xlsx")})
    assert res.status_code == 200
    assert "error" in res.json()
