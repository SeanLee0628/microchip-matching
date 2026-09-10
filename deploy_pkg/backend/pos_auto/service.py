# -*- coding: utf-8 -*-
"""웹 앱에서 POS Report 자동화를 호출하는 얇은 래퍼.

업로드된 바이트를 임시 폴더에 쓰고 CLI 와 동일한 파이프라인(main.run)을 돌린 뒤,
생성된 결과 워크북을 바이트로 돌려준다. 로직은 CLI 와 100% 동일하다.
"""
import argparse
import glob
import shutil
import tempfile
from pathlib import Path
from typing import Optional, Tuple

import openpyxl

from . import main as pos_main

HERE = Path(__file__).resolve().parent
RULES = HERE / "rules.yaml"


def _write(tmp: Path, name: str, data: Optional[bytes]) -> Optional[str]:
    if not data:
        return None
    p = tmp / name
    p.write_bytes(data)
    return str(p)


def _summary_from(path: Path) -> dict:
    """결과 워크북에서 화면에 보여줄 요약을 뽑는다."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {"summary": [], "counts": {}}
    if "Validation_Summary" in wb.sheetnames:
        ws = wb["Validation_Summary"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0] is not None:
                out["summary"].append({"항목": str(r[0]), "값": r[1]})
    for name in ("POS_Report", "Exceptions", "Customer_Match_Review",
                 "Address_Master_Needed", "QTN_Usage_Ledger"):
        if name in wb.sheetnames:
            out["counts"][name] = max(wb[name].max_row - 1, 0)
    wb.close()
    return out


def run_pos(rawdata: bytes, start_date: str, end_date: str,
            crm: Optional[bytes] = None, address_master: Optional[bytes] = None,
            approved_mapping: Optional[bytes] = None,
            mode: str = "draft") -> Tuple[bytes, str, dict]:
    """returns (xlsx_bytes, filename, summary_dict)  |  실패 시 ValueError"""
    tmp = Path(tempfile.mkdtemp(prefix="pos_auto_"))
    try:
        raw_path = _write(tmp, "RAWDATA.xlsx", rawdata)
        if not raw_path:
            raise ValueError("RAWDATA 파일이 필요합니다.")
        args = argparse.Namespace(
            rawdata=raw_path,
            template=None,
            start_date=start_date,
            end_date=end_date,
            output_dir=str(tmp / "out"),
            log_dir=str(tmp / "logs"),
            config=str(RULES),
            approved_mapping=_write(tmp, "approved_mapping.xlsx", approved_mapping),
            address_master=_write(tmp, "address_master.xlsx", address_master),
            crm_customers=_write(tmp, "crm.xlsx", crm),
            mode=mode,
        )
        rc = pos_main.run(args)
        files = sorted(glob.glob(str(tmp / "out" / "Microchip_POS_Result_*.xlsx")))
        if not files:
            if rc == 3:
                raise ValueError("final 모드는 미해결 필수 예외가 0건일 때만 생성됩니다. "
                                 "먼저 검토(draft)로 예외를 해소하세요.")
            raise ValueError("결과 파일이 생성되지 않았습니다. 입력 파일의 시트·컬럼을 확인하세요.")
        out = Path(files[-1])
        return out.read_bytes(), out.name, _summary_from(out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
