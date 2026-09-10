# -*- coding: utf-8 -*-
"""결과 엑셀 작성 — 수식 없이 값만 쓴다."""
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def _sheet(wb: Workbook, title: str, rows: List[dict], columns: List[str] = None):
    ws = wb.create_sheet(title[:31])
    cols = columns or (list(rows[0].keys()) if rows else ["(데이터 없음)"])
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, BOLD, CENTER
    for r, item in enumerate(rows, 2):
        for i, c in enumerate(cols, 1):
            v = item.get(c)
            if isinstance(v, (list, tuple, set)):
                v = ", ".join(str(x) for x in v)
            ws.cell(row=r, column=i, value=v)
    widths = {}
    for i, c in enumerate(cols, 1):
        w = max([len(str(c))] + [len(str(item.get(c, ""))) for item in rows[:200]] or [10])
        widths[i] = min(max(w + 2, 10), 42)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    ws.freeze_panes = "A2"
    return ws


def write_result(path: Path, sheets: Dict[str, List[dict]], pos_columns: List[str],
                 draft: bool) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    order = ["POS_Report", "Processing_Result", "Exceptions", "Customer_Match_Review",
             "Validation_Summary", "QTN_Usage_Ledger", "POS_Field_Mapping",
             "Address_Master_Needed", "Run_Log"]
    for name in order:
        rows = sheets.get(name, [])
        cols = pos_columns if name == "POS_Report" else None
        ws = _sheet(wb, name, rows, cols)
        if name == "POS_Report" and draft:
            ws.sheet_view.showGridLines = True
            ws.oddHeader.center.text = "DRAFT - 승인 전"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_address_master_template(path: Path, rows: List[dict]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    cols = ["고객코드", "거래처명", "영문명(지역포함)",
            "purchasing_city", "purchasing_state", "purchasing_postal", "purchasing_country",
            "end_city", "end_state", "end_postal", "end_country",
            "disti_customer_number", "CRM 주소(참고)", "미입력 필드"]
    _sheet(wb, "customer_address_master", rows, cols)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
