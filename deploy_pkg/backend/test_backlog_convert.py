# -*- coding: utf-8 -*-
"""마이크론 백로그 원본 → Backlog Shipment Report 변환 테스트.

핵심 전제 두 개를 지킨다:
  1) 원본 열 순서·개수가 회차마다 달라도(9/2 15열 / 9/7 18열) 출력 순서는 고정이다.
  2) 키에 후보가 여러 개면 채우지 않는다 — 같은 PO# 에 담당자가 둘인 경우가 실제로 있다.

실행: python -m pytest test_backlog_convert.py -q   (backend/ 에서)
"""
import io
from datetime import datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient

import backlog_convert as bc
import main
import main_aws

BACKENDS = [pytest.param(main, id="main"), pytest.param(main_aws, id="main_aws")]

# 9/7 원본 열 순서 그대로 (SO 가 첫 열, CHANNEL_CODE 가 뒤쪽)
SRC_0907 = ["SO", "ORDER_TYPE", "PLANT", "LINE_ITEM_BLOCK_CODE", "LINE_ITEM_BLOCK_DESC",
            "SAP_NO", "MPN", "DID", "BOX_TYPE", "QTY", "OPEN_ORDER_VALUE",
            "END_CUSTOMER_NAME", "CRD", "MAD", "CUSTOMER_MATERIAL", "CHANNEL_CODE",
            "PURCH_ORDER_NO", "DELIVERY_NUMBER"]
# 9/2 원본 — LINE_ITEM_BLOCK_* 과 DELIVERY_NUMBER 가 없다
SRC_0902 = [c for c in SRC_0907 if not c.startswith("LINE_ITEM") and c != "DELIVERY_NUMBER"]


def _src_row(so, po, mpn, qty=1000, value=21600.0, did="V00H", cust_mat=None):
    return {"SO": so, "ORDER_TYPE": "OR", "PLANT": "SG15", "LINE_ITEM_BLOCK_CODE": None,
            "LINE_ITEM_BLOCK_DESC": None, "SAP_NO": 12345, "MPN": mpn, "DID": did,
            "BOX_TYPE": "TR", "QTY": qty, "OPEN_ORDER_VALUE": value,
            "END_CUSTOMER_NAME": "UNITRON KR", "CRD": datetime(2026, 9, 14),
            "MAD": datetime(2026, 9, 21), "CUSTOMER_MATERIAL": cust_mat,
            "CHANNEL_CODE": 2, "PURCH_ORDER_NO": po, "DELIVERY_NUMBER": None}


def _prev_row(so, po, mpn, dbc=21.6, fse="JH", cust="HANWHA VISION"):
    r = _src_row(so, po, mpn)
    r.update({"DBC": dbc, "FSE": fse, "CUST": cust, "OPEN COST": 21.6})
    return r


def _xlsx(header, rows, sheet="Unitron"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _prev_xlsx(rows):
    header = bc.OUT_COLS
    return _xlsx(header, rows, sheet="Backlog")


# ---------- 헤더 읽기 ----------

def test_reads_sheet_whatever_its_name_and_column_order():
    """원본 시트명이 제각각이고(Sheet1/Unitron) 열 순서도 다르다."""
    for sheet in ("Sheet1", "Unitron", "Backlog", "Disti"):
        h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")], sheet=sheet))
        assert rows and rows[0]["SO"] == "S1", sheet


def test_header_names_are_matched_loosely():
    """'OPEN COST' / 'open_cost' 같은 표기 차이를 같은 열로 본다."""
    assert bc.norm_header("OPEN COST") == bc.norm_header("open_cost") == "OPENCOST"
    assert bc.norm_header(" PURCH_ORDER_NO ") == bc.norm_header("Purch Order No")


def test_non_backlog_file_is_rejected():
    h, rows = bc.read_sheet(_xlsx(["A", "B", "C"], [{"A": 1, "B": 2, "C": 3}]))
    assert h is None and rows is None


def test_blank_rows_are_skipped():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [
        _src_row("S1", "P1", "M1"), {k: None for k in SRC_0907}, _src_row("S2", "P1", "M1")]))
    assert [r["SO"] for r in rows] == ["S1", "S2"]


# ---------- 열 순서 ----------

def test_output_order_is_fixed_regardless_of_source_order():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    res = bc.convert(h, rows)
    assert res["columns"] == bc.OUT_COLS


def test_optional_columns_are_dropped_when_source_lacks_them():
    """9/2 원본처럼 LINE_ITEM_BLOCK_*·DELIVERY_NUMBER 가 없으면 빼고 순서는 유지."""
    h, rows = bc.read_sheet(_xlsx(SRC_0902, [_src_row("S1", "P1", "M1")]))
    res = bc.convert(h, rows)
    assert res["columns"] == [c for c in bc.OUT_COLS if c not in bc.OPTIONAL_COLS]
    assert res["summary"]["dropped_optional"] == [
        "LINE_ITEM_BLOCK_CODE", "LINE_ITEM_BLOCK_DESC", "DELIVERY_NUMBER"]


def test_derived_columns_always_present_even_without_prev():
    h, rows = bc.read_sheet(_xlsx(SRC_0902, [_src_row("S1", "P1", "M1")]))
    res = bc.convert(h, rows, None)
    for c in bc.DERIVED_COLS:
        assert c in res["columns"], c
    assert res["rows"][0]["DBC"] is None and res["rows"][0]["FSE"] is None


# ---------- OPEN COST ----------

def test_open_cost_is_value_over_qty():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1", qty=1000, value=21600.0)]))
    assert bc.convert(h, rows)["rows"][0]["OPEN COST"] == 21.6


def test_open_cost_is_none_when_qty_is_zero():
    """0으로 나누지 않는다."""
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1", qty=0, value=100.0)]))
    assert bc.convert(h, rows)["rows"][0]["OPEN COST"] is None


def test_open_cost_written_as_live_formula():
    """완성본이 =Q2/P2 수식이라 그대로 수식으로 넣는다."""
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    wb = openpyxl.load_workbook(io.BytesIO(bc.build_xlsx_bytes(bc.convert(h, rows))))
    ws = wb.active
    head = [c.value for c in ws[1]]
    col = head.index("OPEN COST") + 1
    assert ws.cell(2, col).value == "=Q2/P2", ws.cell(2, col).value


# ---------- DBC ----------

def test_dbc_prefers_so_over_mpn():
    """같은 MPN 이 다른 DBC 를 가져도 SO# 가 맞으면 그 값을 쓴다."""
    prev = [_prev_row("S1", "P1", "M1", dbc=21.6), _prev_row("S9", "P9", "M1", dbc=99.9)]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["DBC"] == 21.6
    assert res["fills"][0]["DBC"] == "SO#"


def test_dbc_falls_back_to_mpn_for_new_so():
    prev = [_prev_row("OLD", "P1", "M1", dbc=21.6)]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["DBC"] == 21.6
    assert res["fills"][0]["DBC"] == "MPN"


def test_dbc_left_blank_when_mpn_has_two_prices():
    """MPN 하나에 DBC 가 둘이면 찍지 않는다 (실데이터에 4종 있었다)."""
    prev = [_prev_row("A", "PA", "M1", dbc=23.96), _prev_row("B", "PB", "M1", dbc=17.85)]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "PZ", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["DBC"] is None
    note = [r for r in res["review"] if r["field"] == "DBC"][0]
    assert "후보가 여러 개" in note["reason"]
    assert sorted(note["candidates"]) == ["17.85", "23.96"]


def test_dbc_blank_when_mpn_is_new():
    prev = [_prev_row("A", "PA", "M1")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "PZ", "M_NEW")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["DBC"] is None
    assert [r["reason"] for r in res["review"] if r["field"] == "DBC"] == ["이전 백록에 없음"]


# ---------- FSE / CUST ----------

def test_fse_cust_prefer_so_over_po():
    """사용자가 손으로 하던 순서 — 이전 백록에서 맞춰둔 SO# 가 먼저다."""
    prev = [_prev_row("S1", "P1", "M1", fse="KATE", cust="SKYHIGH"),
            _prev_row("S2", "P1", "M1", fse="JH", cust="HANWHA")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert (res["rows"][0]["FSE"], res["rows"][0]["CUST"]) == ("KATE", "SKYHIGH")
    assert res["fills"][0]["FSE"] == "SO#"


def test_fse_cust_fall_back_to_po_for_new_so():
    prev = [_prev_row("OLD", "P1", "M1", fse="KATE", cust="SKYHIGH")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert (res["rows"][0]["FSE"], res["rows"][0]["CUST"]) == ("KATE", "SKYHIGH")
    assert res["fills"][0]["CUST"] == "PO#"


def test_two_owners_on_one_po_is_left_blank_not_guessed():
    """사용자가 지적한 그 경우. VLOOKUP 은 첫 값을 집지만 근거가 없다."""
    prev = [_prev_row("A", "P1", "M1", fse="KATE", cust="SKYHIGH"),
            _prev_row("B", "P1", "M1", fse="JH", cust="HANWHA")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["FSE"] is None and res["rows"][0]["CUST"] is None
    fse = [r for r in res["review"] if r["field"] == "FSE"][0]
    assert "PO#" in fse["reason"] and sorted(fse["candidates"]) == ["JH", "KATE"]


def test_one_field_resolves_even_if_the_other_is_ambiguous():
    """같은 PO# 에 담당자는 둘인데 고객사는 하나면, 고객사는 채운다."""
    prev = [_prev_row("A", "P1", "M1", fse="KATE", cust="SAME"),
            _prev_row("B", "P1", "M1", fse="JH", cust="SAME")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("NEW", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["FSE"] is None
    assert res["rows"][0]["CUST"] == "SAME"


def test_keys_match_despite_type_and_space_noise():
    """SO#·PO# 가 숫자로 읽히거나 공백이 붙어 와도 같은 키로 본다."""
    prev = [_prev_row(611301667, " p1 ", "M1", fse="JH")]
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row(" 611301667 ", "P1", "M1")]))
    res = bc.convert(h, rows, prev)
    assert res["rows"][0]["FSE"] == "JH"


# ---------- 엑셀 ----------

def test_xlsx_matches_finished_file_styling():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    wb = openpyxl.load_workbook(io.BytesIO(bc.build_xlsx_bytes(bc.convert(h, rows), "Unitron")))
    ws = wb["Unitron"]
    head = [c.value for c in ws[1]]
    assert head == bc.OUT_COLS
    assert ws["A1"].fill.start_color.rgb == bc.C_HDR
    assert ws.cell(1, head.index("OPEN COST") + 1).fill.start_color.rgb == bc.C_HDR_CALC
    assert ws.cell(1, head.index("OPEN_ORDER_VALUE") + 1).fill.start_color.rgb == bc.C_HDR_CALC
    assert (ws["A1"].font.name, ws["A1"].font.size, ws["A1"].font.bold) == ("맑은 고딕", 10, True)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:V2"
    assert ws.cell(2, head.index("QTY") + 1).number_format == bc.FMT_QTY
    assert ws.cell(2, head.index("CRD") + 1).number_format == bc.FMT_DATE


def test_sheet_title_keeps_source_name_but_not_sheet1():
    assert bc._sheet_title("Unitron") == "Unitron"
    assert bc._sheet_title("Sheet1") == "Backlog"
    assert bc._sheet_title("") == "Backlog"


def test_review_sheet_only_when_something_needs_checking():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    clean = openpyxl.load_workbook(io.BytesIO(
        bc.build_xlsx_bytes(bc.convert(h, rows, [_prev_row("S1", "P1", "M1")]))))
    assert "확인필요" not in clean.sheetnames
    # 이전 백록은 줬는데 그 안에 이 SO·PO·MPN 이 없는 경우
    dirty = openpyxl.load_workbook(io.BytesIO(
        bc.build_xlsx_bytes(bc.convert(h, rows, [_prev_row("OTHER", "PX", "MX")]))))
    assert "확인필요" in dirty.sheetnames
    assert dirty["확인필요"].max_row == 4, "DBC·FSE·CUST 3건 + 헤더"


def test_no_prev_file_produces_no_review_sheet():
    h, rows = bc.read_sheet(_xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    wb = openpyxl.load_workbook(io.BytesIO(bc.build_xlsx_bytes(bc.convert(h, rows, None))))
    assert "확인필요" not in wb.sheetnames and len(wb.sheetnames) == 1


# ---------- 엔드포인트 ----------

def _post(mod, src, prev=None):
    files = {"file": ("src.xlsx", src, "application/xlsx")}
    if prev is not None:
        files["prev"] = ("prev.xlsx", prev, "application/xlsx")
    res = TestClient(mod.app).post("/api/backlog-convert/preview", files=files)
    assert res.status_code == 200, res.text
    return res.json()


@pytest.mark.parametrize("mod", BACKENDS)
def test_preview_returns_columns_summary_and_review(mod):
    src = _xlsx(SRC_0907, [_src_row("S1", "P1", "M1")])
    body = _post(mod, src, _prev_xlsx([_prev_row("S1", "P1", "M1")]))
    assert body.get("error") is None
    assert body["columns"] == bc.OUT_COLS
    assert body["summary"]["row_count"] == 1
    assert body["summary"]["fill_counts"]["DBC"]["SO#"] == 1
    assert body["review"] == []


@pytest.mark.parametrize("mod", BACKENDS)
def test_preview_without_prev_leaves_blanks_but_no_review_noise(mod):
    """참조 파일을 안 줬으면 비는 게 당연하다 — 확인 필요로 쌓으면 만 건이 된다."""
    body = _post(mod, _xlsx(SRC_0907, [_src_row("S1", "P1", "M1")]))
    assert body["summary"]["has_prev"] is False
    assert body["summary"]["review_count"] == 0
    assert body["review"] == []
    row = body["rows"][0]
    assert row["DBC"] is None and row["FSE"] is None and row["CUST"] is None


@pytest.mark.parametrize("mod", BACKENDS)
def test_preview_caps_rows_but_keeps_full_count(mod):
    """3천 행대라 화면으로는 앞부분만 보낸다."""
    many = [_src_row(f"S{i}", "P1", "M1") for i in range(bc.PREVIEW_ROWS + 50)]
    body = _post(mod, _xlsx(SRC_0907, many))
    assert body["summary"]["row_count"] == bc.PREVIEW_ROWS + 50
    assert len(body["rows"]) == bc.PREVIEW_ROWS


@pytest.mark.parametrize("mod", BACKENDS)
def test_preview_rejects_non_backlog(mod):
    body = _post(mod, _xlsx(["A", "B"], [{"A": 1, "B": 2}]))
    assert "Backlog" in body["error"]


@pytest.mark.parametrize("mod", BACKENDS)
def test_export_returns_xlsx_with_counts_in_headers(mod):
    src = _xlsx(SRC_0907, [_src_row("S1", "P1", "M1"), _src_row("S2", "P1", "M1")])
    prev = _prev_xlsx([_prev_row("S1", "P1", "M1")])
    res = TestClient(mod.app).post("/api/backlog-convert/export", files={
        "file": ("src.xlsx", src, "application/xlsx"),
        "prev": ("prev.xlsx", prev, "application/xlsx")})
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
    assert res.headers["x-row-count"] == "2"
    assert res.headers["x-review-count"] == "0", "S2 는 PO# 로 채워진다"
    res.headers["content-disposition"].encode("latin-1")
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in ws[1]] == bc.OUT_COLS
