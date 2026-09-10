# -*- coding: utf-8 -*-
"""CRD 현황판 파싱·엑셀 내보내기 회귀 테스트 (2026-09-10 영업1실 요청서 반영분).

순수 판정 로직은 test_crd_board.py. 여기서는 엑셀 경계 — 컬럼을 제대로 읽는지,
내보낸 파일이 "원본 전체 행 + 계산열" 인지를 지킨다.
실행: python -m pytest test_crd_board_api.py -q   (backend/ 에서)
"""
import io
from datetime import date, datetime

import openpyxl
import main_aws as m

# 첨부 자료(Backlog Shipment Report - 260909_수정.xlsx)의 헤더 순서 그대로
HDR = ["ORDER_TYPE", "SO", "PURCH_ORDER_NO", "SAP_NO", "CUSTOMER_MATERIAL", "MPN", "DID",
       "END_CUSTOMER_NAME", "CRD", "MAD", "PLANT", "BOX_TYPE", "QTY",
       "OPEN_ORDER_VALUE", "DELIVERY_NUMBER", "DBC", "FSE", "CUST"]


def _row(so, crd, mad, **kw):
    v = {"ORDER_TYPE": "OR", "SO": so, "PURCH_ORDER_NO": "45HWA260203-01", "SAP_NO": "S1",
         "CUSTOMER_MATERIAL": "M320300", "MPN": "MT29F8G08ABBCAH4-IT:C", "DID": "M71M",
         "END_CUSTOMER_NAME": "UNITRON KR", "CRD": crd, "MAD": mad, "PLANT": "SG15",
         "BOX_TYPE": "DRY PACK", "QTY": 1260, "OPEN_ORDER_VALUE": 1000,
         "DELIVERY_NUMBER": None, "DBC": None, "FSE": "JH", "CUST": "HANWHA VISION"}
    v.update(kw)
    return [v[h] for h in HDR]


def _xlsx(rows, hdr=HDR, sheet="Backlog"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(hdr)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(content):
    return openpyxl.load_workbook(io.BytesIO(content)).active


class TestParseColumns:
    """요청서 표에 나가는 열을 전부 읽어야 한다."""

    def test_end_customer_name_is_read(self):
        # 예전 코드는 "End customer" 만 찾아서 고객명이 전 행 빈칸이었다.
        o = m._parse_backlog_orders(_xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7))]))[0]
        assert o["customer"] == "UNITRON KR"

    def test_legacy_end_customer_header_still_works(self):
        hdr = [("End customer" if h == "END_CUSTOMER_NAME" else h) for h in HDR]
        o = m._parse_backlog_orders(
            _xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7))], hdr=hdr))[0]
        assert o["customer"] == "UNITRON KR"

    def test_table_columns_are_read(self):
        o = m._parse_backlog_orders(_xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7))]))[0]
        assert o["po"] == "45HWA260203-01"
        assert o["cust_material"] == "M320300"
        assert o["plant"] == "SG15"
        assert o["box_type"] == "DRY PACK"
        assert o["cust"] == "HANWHA VISION"
        assert o["fse"] == "JH"

    def test_fse_cust_zero_becomes_blank(self):
        # 첨부 자료에 FSE·CUST 가 0 으로 들어온 행이 3건 있다 → 공란 취급
        o = m._parse_backlog_orders(
            _xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7), FSE=0, CUST=0)]))[0]
        assert o["fse"] is None
        assert o["cust"] is None

    def test_missing_fse_cust_columns_are_blank(self):
        hdr = [h for h in HDR if h not in ("FSE", "CUST")]
        rows = [[v for h, v in zip(HDR, _row("S1", date(2026, 7, 20), date(2026, 9, 7)))
                 if h in hdr]]
        o = m._parse_backlog_orders(_xlsx(rows, hdr=hdr))[0]
        assert o["fse"] is None and o["cust"] is None

    def test_delivery_number_marks_shipped(self):
        orders = m._parse_backlog_orders(_xlsx([
            _row("S1", date(2026, 7, 20), date(2026, 9, 7)),
            _row("S2", date(2026, 7, 20), date(2026, 9, 7), DELIVERY_NUMBER="8012345"),
        ]))
        assert [o["open"] for o in orders] == [True, False]


class TestElapsedExport:
    """현황보기 다운로드 = 원본 전체 행 + 경과일수 열 (화면 필터와 무관)."""

    def _export(self, rows):
        hdr, data = m._bl_read_raw_sheet(_xlsx(rows))
        extras = []
        i_crd, i_mad = hdr.index("CRD"), hdr.index("MAD")
        for r in data:
            crd, mad = m._bl_cell_date(r[i_crd]), m._bl_cell_date(r[i_mad])
            extras.append([None if (crd is None or mad is None) else (mad - crd).days])
        wb = m._bl_export_workbook(hdr, data, ["경과일수"], extras)
        buf = io.BytesIO()
        wb.save(buf)
        return _sheet(buf.getvalue())

    def test_elapsed_column_sits_right_after_mad(self):
        ws = self._export([_row("S1", date(2026, 7, 20), date(2026, 9, 7))])
        head = [c.value for c in ws[1]]
        assert head[head.index("MAD") + 1] == "경과일수"

    def test_elapsed_value_is_mad_minus_crd(self):
        # 요청서 예시 R11: CRD 2026-07-20, MAD 2026-09-07 → 49
        ws = self._export([_row("S1", date(2026, 7, 20), date(2026, 9, 7))])
        head = [c.value for c in ws[1]]
        assert ws.cell(2, head.index("경과일수") + 1).value == 49

    def test_all_rows_exported_including_not_elapsed(self):
        rows = [_row("S1", date(2026, 7, 20), date(2026, 9, 7)),    # 경과 49
                _row("S2", date(2026, 9, 9), date(2026, 9, 9)),     # 경과 0
                _row("S3", date(2026, 9, 20), date(2026, 9, 9))]    # 경과 -11
        ws = self._export(rows)
        head = [c.value for c in ws[1]]
        col = head.index("경과일수") + 1
        assert ws.max_row == 4                                       # 헤더 + 3행 전부
        assert [ws.cell(r, col).value for r in (2, 3, 4)] == [49, 0, -11]

    def test_original_columns_preserved(self):
        ws = self._export([_row("S1", date(2026, 7, 20), date(2026, 9, 7))])
        head = [c.value for c in ws[1]]
        assert [h for h in head if h != "경과일수"] == HDR

    def test_missing_date_is_na(self):
        ws = self._export([_row("S1", None, date(2026, 9, 7))])
        head = [c.value for c in ws[1]]
        assert ws.cell(2, head.index("경과일수") + 1).value == "N/A"


class TestCompareExport:
    """변화보기 다운로드 = 현재 백록 전체 행 + 이전 MAD·GAP 열."""

    def _export(self, prev_rows, cur_rows):
        prev_orders = m._parse_backlog_orders(_xlsx(prev_rows))
        hdr, data = m._bl_read_raw_sheet(_xlsx(cur_rows))
        pmad = {o["so"]: o["mad"] for o in prev_orders if o.get("so")}
        i_so, i_mad = hdr.index("SO"), hdr.index("MAD")
        vals = []
        for r in data:
            so = m._bl_clean(r[i_so])
            mad = m._bl_cell_date(r[i_mad])
            p = pmad.get(so)
            vals.append([p, None] if (p is None or mad is None) else [p, (mad - p).days])
        wb = m._bl_export_workbook(hdr, data, ["이전 MAD", "GAP"], vals)
        buf = io.BytesIO()
        wb.save(buf)
        ws = _sheet(buf.getvalue())
        head = [c.value for c in ws[1]]
        return ws, head

    def test_columns_inserted_after_mad(self):
        ws, head = self._export([_row("S1", date(2026, 9, 11), date(2026, 9, 14))],
                                [_row("S1", date(2026, 9, 11), date(2026, 9, 11))])
        i = head.index("MAD")
        assert head[i + 1:i + 3] == ["이전 MAD", "GAP"]

    def test_gap_is_current_minus_previous(self):
        # 요청서 예시 R11: 9/9 MAD 09-11, 8/31 MAD 09-14 → GAP -3
        ws, head = self._export([_row("S1", date(2026, 9, 11), date(2026, 9, 14))],
                                [_row("S1", date(2026, 9, 11), date(2026, 9, 11))])
        assert ws.cell(2, head.index("GAP") + 1).value == -3

    def test_push_out_gap_positive(self):
        # 요청서 예시 R20: 9/9 MAD 09-14, 8/31 MAD 09-07 → GAP +7
        ws, head = self._export([_row("S1", date(2026, 6, 11), date(2026, 9, 7))],
                                [_row("S1", date(2026, 6, 11), date(2026, 9, 14))])
        assert ws.cell(2, head.index("GAP") + 1).value == 7

    def test_new_so_is_na(self):
        ws, head = self._export([_row("OLD", date(2026, 9, 1), date(2026, 9, 1))],
                                [_row("NEW", date(2026, 9, 14), date(2026, 9, 14))])
        assert ws.cell(2, head.index("이전 MAD") + 1).value == "N/A"
        assert ws.cell(2, head.index("GAP") + 1).value == "N/A"

    def test_unchanged_rows_still_in_file(self):
        """화면은 변화된 것만, 파일은 전체 백록 라인 — 요청서 N7 코멘트."""
        ws, head = self._export(
            [_row("S1", date(2026, 9, 1), date(2026, 9, 1)),
             _row("S2", date(2026, 9, 1), date(2026, 9, 1))],
            [_row("S1", date(2026, 9, 1), date(2026, 9, 1)),      # 변화 없음
             _row("S2", date(2026, 9, 1), date(2026, 9, 8))])     # +7
        assert ws.max_row == 3
        col = head.index("GAP") + 1
        assert [ws.cell(r, col).value for r in (2, 3)] == [0, 7]


class TestEndpoints:
    """업로드 → 응답/다운로드 실제 왕복."""

    def _client(self):
        from fastapi.testclient import TestClient
        return TestClient(m.app)

    def _up(self, content, name="bl.xlsx"):
        return {"file": (name, content,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    def test_board_returns_table_fields_and_elapsed_count(self):
        content = _xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7)),
                         _row("S2", date(2026, 9, 9), date(2026, 9, 9))])
        r = self._client().post("/api/crd-board", files=self._up(content))
        d = r.json()
        assert d.get("error") is None
        assert d["open_count"] == 2
        assert d["elapsed_count"] == 1          # 경과일수 0 은 경과 아님
        row = d["board"][0]
        for k in ("po", "cust_material", "plant", "box_type", "cust", "delay_days", "elapsed"):
            assert k in row
        assert row["customer"] == "UNITRON KR"

    def test_board_sorted_by_mad(self):
        content = _xlsx([_row("LATE", date(2026, 1, 1), date(2026, 9, 20)),
                         _row("SOON", date(2026, 1, 1), date(2026, 9, 7))])
        d = self._client().post("/api/crd-board", files=self._up(content)).json()
        assert [c["so"] for c in d["board"]] == ["SOON", "LATE"]

    def test_export_returns_xlsx_with_all_rows(self):
        content = _xlsx([_row("S1", date(2026, 7, 20), date(2026, 9, 7)),
                         _row("S2", date(2026, 9, 9), date(2026, 9, 9))])
        r = self._client().post("/api/crd-board/export", files=self._up(content))
        assert r.status_code == 200
        assert r.headers["x-row-count"] == "2"
        assert r.headers["x-elapsed-count"] == "1"
        ws = _sheet(r.content)
        assert ws.max_row == 3
        assert "경과일수" in [c.value for c in ws[1]]

    def test_compare_returns_changed_both_directions(self):
        prev = _xlsx([_row("S1", date(2026, 9, 1), date(2026, 9, 1)),
                      _row("S2", date(2026, 9, 1), date(2026, 9, 8)),
                      _row("S3", date(2026, 9, 1), date(2026, 9, 1))])
        cur = _xlsx([_row("S1", date(2026, 9, 1), date(2026, 9, 8)),    # +7 밀림
                     _row("S2", date(2026, 9, 1), date(2026, 9, 1)),    # -7 당겨짐
                     _row("S3", date(2026, 9, 1), date(2026, 9, 1)),    # 동일
                     _row("S4", date(2026, 9, 1), date(2026, 9, 1))])   # 신규
        c = self._client()
        r = c.post("/api/crd-board/compare", files={
            "prev": ("p.xlsx", prev, "application/octet-stream"),
            "current": ("c.xlsx", cur, "application/octet-stream")})
        d = r.json()
        assert d["summary"]["changed"] == 3          # 밀림 + 당겨짐 + 신규
        assert d["summary"]["slipped"] == 1
        assert d["summary"]["improved"] == 1
        assert d["summary"]["same"] == 1
        assert d["summary"]["new"] == 1
        gaps = {x["so"]: x["gap_days"] for x in d["changed"]}
        assert gaps == {"S1": 7, "S2": -7, "S4": None}

    def test_compare_export_returns_xlsx(self):
        prev = _xlsx([_row("S1", date(2026, 9, 1), date(2026, 9, 1))])
        cur = _xlsx([_row("S1", date(2026, 9, 1), date(2026, 9, 8)),
                     _row("S2", date(2026, 9, 1), date(2026, 9, 1))])
        r = self._client().post("/api/crd-board/compare/export", files={
            "prev": ("p.xlsx", prev, "application/octet-stream"),
            "current": ("c.xlsx", cur, "application/octet-stream")})
        assert r.status_code == 200
        assert r.headers["x-row-count"] == "2"
        assert r.headers["x-push-out"] == "1"
        assert r.headers["x-na-count"] == "1"        # S2 는 이전 백록에 없음
        ws = _sheet(r.content)
        head = [c.value for c in ws[1]]
        assert "이전 MAD" in head and "GAP" in head

    def test_non_backlog_file_rejected(self):
        content = _xlsx([[1, 2, 3]], hdr=["A", "B", "C"])
        d = self._client().post("/api/crd-board", files=self._up(content)).json()
        assert "Backlog Shipment Report" in d["error"]
