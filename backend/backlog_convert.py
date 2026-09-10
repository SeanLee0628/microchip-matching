"""마이크론 백로그 원본 → Backlog Shipment Report 재배열 + DBC·FSE·CUST 자동 채움.

main.py(로컬) · main_aws.py(배포) 공용. 사본을 따로 두면 반드시 갈라진다.

마이크론 담당자가 보내주는 원본은 회차마다 열 순서·개수가 다르다
(9/2 원본 15열 / 9/7 원본 18열 — LINE_ITEM_BLOCK_*·DELIVERY_NUMBER 유무).
그래서 위치가 아니라 **헤더 이름**으로 찾아 정해진 순서로 다시 세운다.

DBC·FSE·CUST 는 원본에 없다. 이전 회차 완성본(이전 백록)을 참조표로 써서 채운다.
사용자가 VLOOKUP 으로 하던 순서를 그대로 옮긴 것:

  DBC      : SO# → MPN
  FSE·CUST : SO# → PO#

실측(첨부 9/2 완성본 → 9/7 원본, 3,260행):
  DBC      SO# 로 3,227행 채움 · 실제 값과 **불일치 0** / MPN 폴백 27행(전부 정확) / 확인필요 6행
  FSE·CUST SO# 로 3,227행 · 신규 SO 33행은 PO# 로 전부 유일하게 결정 → 확인필요 0행

**같은 키에 값이 둘 이상이면 추측하지 않는다.** 비워 두고 '확인 필요' 로 뽑는다.
사용자가 지적한 "같은 PO# 인데 담당자·고객사가 다른 경우"가 실제로 있다
(9/7 기준 PO# 400개 중 8개, 50행). VLOOKUP 은 첫 번째 값을 집어오지만 그건 근거가 없다.

단위테스트: test_backlog_convert.py
"""
import io
import re
from datetime import date, datetime

# 출력 열 순서 — 사용자 지정. 이 순서는 고정이다.
OUT_COLS = [
    "ORDER_TYPE", "CHANNEL_CODE", "SO", "LINE_ITEM_BLOCK_CODE", "LINE_ITEM_BLOCK_DESC",
    "PURCH_ORDER_NO", "SAP_NO", "CUSTOMER_MATERIAL", "MPN", "DID", "END_CUSTOMER_NAME",
    "CRD", "MAD", "PLANT", "BOX_TYPE", "QTY", "OPEN_ORDER_VALUE", "OPEN COST",
    "DELIVERY_NUMBER", "DBC", "FSE", "CUST",
]

# 원본에 없으면 빼는 열 (사용자 확인: "생략 가능 / 순서 문제없음").
# DELIVERY_NUMBER 도 여기 둔다 — 9/2 원본에 없어서 9/2 완성본에도 없다.
OPTIONAL_COLS = {"LINE_ITEM_BLOCK_CODE", "LINE_ITEM_BLOCK_DESC", "DELIVERY_NUMBER"}

# 원본에 없고 앱이 만드는 열. 항상 붙는다.
DERIVED_COLS = ["OPEN COST", "DBC", "FSE", "CUST"]

# 이 열들이 다 있는 시트를 원본 시트로 본다 (시트명이 Sheet1/Unitron/Backlog 로 제각각)
REQUIRED_COLS = {"SO", "MPN", "QTY", "CRD", "MAD"}

FILL_KEYS = ("DBC", "FSE", "CUST")

# 화면으로는 앞부분만 보낸다 (실파일이 3천 행대라 전체를 JSON 으로 실으면 무겁다).
# 엑셀 내려받기는 파일을 다시 읽어 전수로 만든다.
PREVIEW_ROWS = 300

# 완성본에서 그대로 옮긴 서식
C_HDR = "FF0077C8"        # 헤더 파랑
C_HDR_CALC = "FFFFFF00"   # OPEN_ORDER_VALUE·OPEN COST 는 노랑
FMT_DATE = "mm-dd-yy"
FMT_QTY = r'_-* #,##0_-;\-* #,##0_-;_-* "-"_-;_-@_-'
FMT_DBC = r'_-* #,##0.00_-;\-* #,##0.00_-;_-* "-"_-;_-@_-'
WIDTHS = {
    "CHANNEL_CODE": 9.2, "SO": 14.1, "PURCH_ORDER_NO": 19.5, "SAP_NO": 9.2,
    "MPN": 27.9, "END_CUSTOMER_NAME": 14.6, "CRD": 10.1, "QTY": 9.2,
    "OPEN_ORDER_VALUE": 10.4, "OPEN COST": 11.2, "DELIVERY_NUMBER": 16.1, "DBC": 9.0,
}


# ---------- 헤더 ----------

def norm_header(h):
    """헤더 이름 정규화. 공백·밑줄·대소문자 차이를 무시한다.

    'OPEN COST' / 'OPEN_COST' / 'open cost' 를 같은 열로 본다.
    """
    return re.sub(r"[\s_]+", "", str(h if h is not None else "")).upper()


NORM_OUT = {norm_header(c): c for c in OUT_COLS}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return v


def _key(v):
    """매칭 키 정규화. SO#·PO#·MPN 이 숫자로 읽히거나 공백이 붙어 와도 같은 키로 묶인다."""
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip().upper()
    return s or None


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return None


# ---------- 읽기 ----------

def read_sheet(contents):
    """엑셀에서 백로그 시트를 찾아 (원본 헤더 리스트, 행 dict 리스트) 로 읽는다.

    행 dict 의 키는 정규화 이름(OUT_COLS 표기), 값은 셀 값 그대로다.
    원본 표기는 따로 돌려준다 — 타이틀은 원본을 따라가도 된다고 하셨다.
    백로그 형식이 아니면 (None, None).
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            for _ in range(10):                      # 헤더가 1행이 아닐 수도 있다
                try:
                    raw = next(rows)
                except StopIteration:
                    break
                labels = [v for v in raw]
                names = [NORM_OUT.get(norm_header(v)) for v in labels]
                if not REQUIRED_COLS.issubset({n for n in names if n}):
                    continue
                out = []
                for r in rows:
                    if all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
                        continue
                    rec = {}
                    for name, v in zip(names, r):
                        if name:
                            rec[name] = _clean(v)
                    out.append(rec)
                header = [str(v).strip() for v in labels if v is not None]
                return header, out
    finally:
        wb.close()
    return None, None


def find_sheet_title(contents):
    """원본에서 백로그 시트의 이름. 출력 시트명을 원본에 맞추려고 쓴다."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            for _ in range(10):
                try:
                    raw = next(rows)
                except StopIteration:
                    break
                names = {NORM_OUT.get(norm_header(v)) for v in raw}
                if REQUIRED_COLS.issubset({n for n in names if n}):
                    return ws.title
    finally:
        wb.close()
    return None


# ---------- 참조표 ----------

def build_reference(prev_rows):
    """이전 백록 → 참조표. 후보가 여러 개인 키는 집합을 그대로 남긴다(추측 금지).

    반환: {"dbc_by_so", "dbc_by_mpn", "fse_by_so", "fse_by_po", "cust_by_so", "cust_by_po"}
          각각 {키: set(후보값)}
    """
    ref = {k: {} for k in ("dbc_by_so", "dbc_by_mpn",
                           "fse_by_so", "fse_by_po", "cust_by_so", "cust_by_po")}
    if not prev_rows:
        return ref

    def put(table, key, val):
        if key is None or val is None:
            return
        ref[table].setdefault(key, set()).add(val)

    for r in prev_rows or []:
        so, po, mpn = _key(r.get("SO")), _key(r.get("PURCH_ORDER_NO")), _key(r.get("MPN"))
        dbc = _num(r.get("DBC"))
        if dbc is not None:
            dbc = round(dbc, 4)
        fse, cust = _clean(r.get("FSE")), _clean(r.get("CUST"))
        put("dbc_by_so", so, dbc)
        put("dbc_by_mpn", mpn, dbc)
        put("fse_by_so", so, fse)
        put("fse_by_po", po, fse)
        put("cust_by_so", so, cust)
        put("cust_by_po", po, cust)
    return ref


def _lookup(table, key):
    """(값, 상태). 상태 = "hit" | "none"(키 없음) | "many"(후보 여러 개).

    후보가 여러 개면 값을 돌려주지 않는다 — VLOOKUP 은 첫 값을 집지만 근거가 없다.
    """
    if key is None or key not in table:
        return None, "none"
    cands = table[key]
    if len(cands) == 1:
        return next(iter(cands)), "hit"
    return None, "many"


def _fill(ref, row):
    """DBC·FSE·CUST 를 SO# → (MPN | PO#) 순서로 채운다.

    반환: ({필드: 값}, {필드: 출처}, [(필드, 사유, 후보)])
      출처 = "SO#" | "MPN" | "PO#" | "" (못 채움)
    """
    so, po, mpn = _key(row.get("SO")), _key(row.get("PURCH_ORDER_NO")), _key(row.get("MPN"))
    plan = {
        "DBC": [("SO#", "dbc_by_so", so), ("MPN", "dbc_by_mpn", mpn)],
        "FSE": [("SO#", "fse_by_so", so), ("PO#", "fse_by_po", po)],
        "CUST": [("SO#", "cust_by_so", so), ("PO#", "cust_by_po", po)],
    }
    values, sources, review = {}, {}, []
    for field, steps in plan.items():
        val, src, many = None, "", []
        for label, table, key in steps:
            v, st = _lookup(ref[table], key)
            if st == "hit":
                val, src = v, label
                break
            if st == "many":
                many.append((label, sorted(str(x) for x in ref[table][key])))
        values[field] = val
        sources[field] = src
        if val is None:
            if many:
                label, cands = many[0]
                review.append((field, f"{label} 에 후보가 여러 개 — 수동 확인", cands))
            else:
                review.append((field, "이전 백록에 없음", []))
    return values, sources, review


# ---------- 변환 ----------

def convert(src_header, src_rows, prev_rows=None):
    """원본 행 → 출력 열 순서로 재배열하고 파생 열을 채운다.

    반환:
      columns : 실제 출력 열 (OUT_COLS 순서, 원본에 없는 선택 열은 빠진다)
      rows    : [{열: 값}] — 'OPEN COST' 는 계산값(엑셀에는 수식으로 들어간다)
      fills   : [{DBC/FSE/CUST: 출처}] rows 와 같은 길이
      review  : 확인 필요 목록
      summary : 열·행 수와 출처별 건수
    """
    present = {c for r in src_rows or [] for c in r}
    columns = [c for c in OUT_COLS
               if c in DERIVED_COLS or c in present or c not in OPTIONAL_COLS]

    ref = build_reference(prev_rows)
    rows, fills, review = [], [], []
    counts = {f: {"SO#": 0, "MPN": 0, "PO#": 0, "": 0} for f in FILL_KEYS}

    for i, r in enumerate(src_rows or [], start=1):
        values, sources, notes = _fill(ref, r)
        out = {}
        for c in columns:
            if c == "OPEN COST":
                v, q = _num(r.get("OPEN_ORDER_VALUE")), _num(r.get("QTY"))
                out[c] = round(v / q, 6) if (v is not None and q) else None
            elif c in FILL_KEYS:
                out[c] = values[c]
            else:
                out[c] = r.get(c)
        rows.append(out)
        fills.append(sources)
        for f in FILL_KEYS:
            counts[f][sources[f]] = counts[f].get(sources[f], 0) + 1
        for field, reason, cands in notes:
            review.append({
                "row": i, "field": field, "reason": reason, "candidates": cands,
                "SO": r.get("SO"), "PURCH_ORDER_NO": r.get("PURCH_ORDER_NO"),
                "MPN": r.get("MPN"), "DID": r.get("DID"),
                "END_CUSTOMER_NAME": r.get("END_CUSTOMER_NAME"),
            })

    return {
        "columns": columns,
        "rows": rows,
        "fills": fills,
        "review": review,
        "summary": {
            "row_count": len(rows),
            "source_columns": len(src_header or []),
            "output_columns": len(columns),
            "dropped_optional": [c for c in OUT_COLS if c in OPTIONAL_COLS and c not in columns],
            "fill_counts": counts,
            "review_count": len(review),
            "has_prev": bool(prev_rows),
        },
    }


# ---------- 엑셀 ----------

def _sheet_title(src_title):
    """원본 시트명을 살린다. 'Sheet1' 같은 기본 이름이면 'Backlog'."""
    t = (src_title or "").strip()
    return "Backlog" if not t or re.fullmatch(r"Sheet\d*", t, re.I) else t[:31]


def build_xlsx_bytes(result, sheet_title=None):
    """완성본과 같은 모양의 xlsx. 헤더색·서식·고정창·필터를 첨부 완성본에서 그대로 옮겼다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns, rows = result["columns"], result["rows"]
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(sheet_title)

    hdr_font = Font(name="맑은 고딕", size=10, bold=True)
    body = Font(name="맑은 고딕", size=10)

    def fill(rgb):
        return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")

    for j, c in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = hdr_font
        cell.fill = fill(C_HDR_CALC if c in ("OPEN_ORDER_VALUE", "OPEN COST") else C_HDR)

    idx = {c: j for j, c in enumerate(columns, start=1)}
    for i, r in enumerate(rows, start=2):
        for c in columns:
            j = idx[c]
            if c == "OPEN COST":
                # 완성본과 같이 살아 있는 수식으로 넣는다 (=OPEN_ORDER_VALUE/QTY)
                v = (f"={get_column_letter(idx['OPEN_ORDER_VALUE'])}{i}"
                     f"/{get_column_letter(idx['QTY'])}{i}"
                     if "OPEN_ORDER_VALUE" in idx and "QTY" in idx else None)
            else:
                v = r.get(c)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = body
            if c in ("CRD", "MAD") and isinstance(v, (datetime, date)):
                cell.number_format = FMT_DATE
            elif c == "QTY":
                cell.number_format = FMT_QTY
            elif c == "DBC":
                cell.number_format = FMT_DBC

    for c, w in WIDTHS.items():
        if c in idx:
            ws.column_dimensions[get_column_letter(idx[c])].width = w
    ws.freeze_panes = "A2"
    if columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"

    review = result.get("review") or []
    if review:
        ws2 = wb.create_sheet("확인필요")
        cols2 = ["행", "항목", "사유", "후보", "SO", "PURCH_ORDER_NO", "MPN", "DID",
                 "END_CUSTOMER_NAME"]
        for j, h in enumerate(cols2, start=1):
            c = ws2.cell(row=1, column=j, value=h)
            c.font, c.fill = hdr_font, fill(C_HDR_CALC)
        for i, rv in enumerate(review, start=2):
            vals = [rv["row"], rv["field"], rv["reason"], ", ".join(rv["candidates"]),
                    rv.get("SO"), rv.get("PURCH_ORDER_NO"), rv.get("MPN"), rv.get("DID"),
                    rv.get("END_CUSTOMER_NAME")]
            for j, v in enumerate(vals, start=1):
                ws2.cell(row=i, column=j, value=v).font = body
        for letter, w in zip("ABCDEFGHI", [7, 8, 30, 22, 14, 19.5, 27.9, 8, 16]):
            ws2.column_dimensions[letter].width = w
        ws2.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
