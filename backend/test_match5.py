# -*- coding: utf-8 -*-
"""영업5실 매칭 순수 로직 테스트 (stdlib unittest, openpyxl).
실행: python -m unittest test_match5 -v   (backend/ 에서)
"""
import io
import unittest
from datetime import date, datetime

import openpyxl

import match5 as m5


class TestMix(unittest.TestCase):
    def test_make_mix_concatenates_code_and_part(self):
        self.assertEqual(m5.make_mix("133742", "24FC01T-E/ST36KVAO"),
                         "13374224FC01T-E/ST36KVAO")

    def test_make_mix_strips_internal_whitespace_and_newlines(self):
        # FCST 일부 MPN에 선행 \n 이 섞여 있음 → 동일 키로 매칭되어야 함
        self.assertEqual(m5.make_mix("133742", "\nMCP1322T-27LE/OTVAO"),
                         m5.make_mix("133742", "MCP1322T-27LE/OTVAO"))

    def test_make_mix_uppercases_part(self):
        self.assertEqual(m5.make_mix("100", "abc-1"), "100ABC-1")

    def test_make_mix_float_code_becomes_int_string(self):
        self.assertEqual(m5.make_mix(131112.0, "23K256T-I/SN"),
                         "13111223K256T-I/SN")

    def test_make_mix_none_returns_none(self):
        self.assertIsNone(m5.make_mix(None, "X"))
        self.assertIsNone(m5.make_mix("100", None))
        self.assertIsNone(m5.make_mix("100", "   "))

    def test_columns_shapes(self):
        self.assertEqual(len(m5.COLUMNS), 27)
        self.assertEqual(m5.COLUMNS[1], "믹스#")
        # 대시보드는 월별/추이/wBL 을 숨긴다
        self.assertNotIn("6월", m5.DASHBOARD_COLUMNS)
        self.assertNotIn("23~25추이", m5.DASHBOARD_COLUMNS)
        self.assertIn("BLOG TTL", m5.DASHBOARD_COLUMNS)


class TestFindHeaderRow(unittest.TestCase):
    def _ws(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        return ws

    def test_header_on_first_row(self):
        ws = self._ws([["PART#", "Qty Due"], ["A", 1]])
        idx, colmap = m5._find_header_row(ws, {"PART#"}, max_scan=5)
        self.assertEqual(idx, 1)
        self.assertEqual(colmap["PART#"], 1)

    def test_header_on_third_row(self):
        ws = self._ws([["title"], [None], ["담당자", "MPN", "Demand Total"], ["a", "b", 3]])
        idx, colmap = m5._find_header_row(ws, {"Demand Total", "MPN"}, max_scan=6)
        self.assertEqual(idx, 3)
        self.assertEqual(colmap["MPN"], 2)
        self.assertEqual(colmap["Demand Total"], 3)

    def test_returns_none_when_not_found(self):
        ws = self._ws([["x", "y"], ["a", "b"]])
        idx, colmap = m5._find_header_row(ws, {"PART#"}, max_scan=5)
        self.assertIsNone(idx)


def _wb_bytes(header_rows, data_rows, sheet_title="Sheet1"):
    """헤더 여러 줄 + 데이터 줄을 가진 xlsx 를 bytes 로."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for row in header_rows:
        ws.append(row)
    for row in data_rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseBlog(unittest.TestCase):
    HEADER = ["PART#", "더존코드", "Qty Due", "PDD", "Change Window",
              "Lead Time Weeks", "Full Manufacturing Cycle Time Weeks", "더존업체명"]

    def _rows(self, rows):
        return _wb_bytes([self.HEADER], rows, sheet_title="260609")

    def test_blog_ttl_and_monthly_bucket_by_pdd(self):
        data = [
            ["23K256T-I/SN", 131112, 3300, datetime(2026, 6, 7), 45, 9, 17, "(주)에이텍"],
            ["23K256T-I/SN", 131112, 6600, datetime(2026, 8, 7), 45, 9, 17, "(주)에이텍"],
        ]
        out = m5.parse_blog(self._rows(data))
        mix = m5.make_mix(131112, "23K256T-I/SN")
        self.assertEqual(out[mix]["blog_ttl"], 9900)
        self.assertEqual(out[mix]["monthly"][6], 3300)
        self.assertEqual(out[mix]["monthly"][8], 6600)

    def test_lead_time_is_max_of_two_columns(self):
        data = [["P1", 100, 10, datetime(2026, 6, 1), 45, 9, 17, "C"]]
        out = m5.parse_blog(self._rows(data))
        self.assertEqual(out[m5.make_mix(100, "P1")]["lead_time"], 17)

    def test_cancel_window_is_earliest_pdd_minus_change_window(self):
        # PDD - Change Window(일). 두 행 중 더 이른 날짜.
        data = [
            ["P1", 100, 10, datetime(2026, 8, 7), 45, 9, 17, "C"],   # 8/7 - 45 = 6/23
            ["P1", 100, 10, datetime(2026, 6, 26), 45, 9, 17, "C"],  # 6/26 - 45 = 5/12 (더 이름)
        ]
        out = m5.parse_blog(self._rows(data))
        self.assertEqual(out[m5.make_mix(100, "P1")]["cancel_window"], date(2026, 5, 12))

    def test_company_and_part_carried(self):
        data = [["P1", 100, 10, datetime(2026, 6, 1), 45, 9, 17, "(주)테스트"]]
        out = m5.parse_blog(self._rows(data))
        rec = out[m5.make_mix(100, "P1")]
        self.assertEqual(rec["더존업체명"], "(주)테스트")
        self.assertEqual(rec["part"], "P1")
        self.assertEqual(rec["더존코드"], "100")


class TestParseShipment(unittest.TestCase):
    HEADER = ["출고일자", "출고번호", "No", "고객코드", "고객", "담당자",
              "품번", "출고수량", "단가"]

    def _rows(self, rows):
        return _wb_bytes([self.HEADER], rows)

    def test_sums_2026_qty_by_mix(self):
        data = [
            [datetime(2026, 1, 2), "IS1", 1, 131149, "(주)에스엔아이", "신성일", "TC4452VAT", 1600, 1],
            [datetime(2026, 3, 5), "IS2", 1, 131149, "(주)에스엔아이", "신성일", "TC4452VAT", 400, 1],
        ]
        out = m5.parse_shipment(self._rows(data))
        mix = m5.make_mix(131149, "TC4452VAT")
        self.assertEqual(out[mix]["y2026"], 2000)
        self.assertEqual(out[mix]["담당자"], "신성일")
        self.assertEqual(out[mix]["고객"], "(주)에스엔아이")
        self.assertEqual(out[mix]["고객코드"], "131149")
        self.assertEqual(out[mix]["part"], "TC4452VAT")

    def test_non_2026_excluded(self):
        data = [[datetime(2025, 12, 31), "X", 1, 100, "C", "S", "P1", 999, 1]]
        out = m5.parse_shipment(self._rows(data))
        self.assertEqual(out[m5.make_mix(100, "P1")]["y2026"], 0)

    def test_cutoff_excludes_on_or_after(self):
        # cutoff_date 지정 시 그 날짜 이상은 제외 (전월 말까지 누적 옵션)
        data = [
            [datetime(2026, 5, 10), "A", 1, 100, "C", "S", "P1", 10, 1],
            [datetime(2026, 6, 20), "B", 1, 100, "C", "S", "P1", 20, 1],
        ]
        out = m5.parse_shipment(self._rows(data), cutoff_date=date(2026, 6, 1))
        self.assertEqual(out[m5.make_mix(100, "P1")]["y2026"], 10)


class TestParseFcst(unittest.TestCase):
    def _bytes(self, data_rows):
        # 헤더 3행: 1·2행 더미, 3행 실제 헤더. 시트명 'Sales Revenue'.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Revenue"
        ws.append(["month band"])
        ws.append(["sub totals"])
        ws.append(["Customer", "MPN", "Demand Total", "Customer Code"])
        for row in data_rows:
            ws.append(row)
        buf = io.BytesIO(); wb.save(buf)
        return buf.getvalue()

    def test_demand_total_keyed_by_mix(self):
        data = [
            ["(주)다나와", "ATMEGA169P-16AU", 6120, 131150],
            ["(주)다나와", "PIC16F1947-I/PT", 3200, 131150],
        ]
        out = m5.parse_fcst(self._bytes(data))
        self.assertEqual(out[m5.make_mix(131150, "ATMEGA169P-16AU")], 6120)
        self.assertEqual(out[m5.make_mix(131150, "PIC16F1947-I/PT")], 3200)

    def test_mpn_with_newline_normalized(self):
        data = [["(주)x", "\nMCP1322T-27LE/OTVAO", 50, 133742]]
        out = m5.parse_fcst(self._bytes(data))
        self.assertEqual(out[m5.make_mix(133742, "MCP1322T-27LE/OTVAO")], 50)

    def test_sums_duplicates(self):
        data = [
            ["c", "P1", 10, 100],
            ["c", "P1", 5, 100],
        ]
        out = m5.parse_fcst(self._bytes(data))
        self.assertEqual(out[m5.make_mix(100, "P1")], 15)


class TestSumAvailableQty(unittest.TestCase):
    def _ws(self, data_rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Jun inventory"
        ws.append(["title row"])  # 1행 더미
        ws.append(["Central", "Sales team", "VENDER", "SR#", "FAMILY", "DID#",
                   "품명", "Part#", "MOBIS ID", "unit", "site", "MOQ", "Package",
                   "FAB", "Q'ty", "SALES", "CUSTOMER", "CRD", "booking",
                   "available Q'ty"])  # 2행 헤더 (Part#=8열, available Q'ty=20열)
        for row in data_rows:
            ws.append(row)
        return ws

    def _row(self, part, oqty, avail):
        return ["A", "5실", "MICROCHIP", ".", ".", ".", "UT", part, ".", "EA",
                ".", ".", ".", ".", oqty, ".", ".", ".", ".", avail]

    def test_sums_available_qty_per_part(self):
        ws = self._ws([
            self._row("PIC16F1947-I/PT", 6400, 6400),
            self._row("PIC16F1947-I/PT", 0, 0),
            self._row("23K256T-I/SN", 3300, 1000),
        ])
        out = m5._sum_available_qty(ws)
        self.assertEqual(out[m5._norm_part("PIC16F1947-I/PT")], 6400)
        self.assertEqual(out[m5._norm_part("23K256T-I/SN")], 1000)

    def test_missing_header_returns_empty(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["nope", "data"])
        self.assertEqual(m5._sum_available_qty(ws), {})


import os
_INV_PATH = r"C:\Users\user\Downloads\cowork\Uniquant_마이크로칩 매칭\Jun_2026_daily warehouse inventory list(영업5실).xlsx"


class TestParseInventoryRealFile(unittest.TestCase):
    @unittest.skipUnless(os.path.exists(_INV_PATH), "real inventory file not present")
    def test_decrypts_and_aggregates(self):
        with open(_INV_PATH, "rb") as f:
            out = m5.parse_inventory(f.read(), password="9178")
        self.assertGreater(len(out), 0)
        # PIC16F1947-I/PT 가 존재하고 합계가 0 이상
        self.assertIn(m5._norm_part("PIC16F1947-I/PT"), out)

    @unittest.skipUnless(os.path.exists(_INV_PATH), "real inventory file not present")
    def test_wrong_password_raises(self):
        with open(_INV_PATH, "rb") as f:
            data = f.read()
        with self.assertRaises(ValueError):
            m5.parse_inventory(data, password="0000")


class TestBuildRecords(unittest.TestCase):
    def test_joins_all_sources_and_computes_balance(self):
        mix = m5.make_mix(131112, "23K256T-I/SN")
        part = m5._norm_part("23K256T-I/SN")
        inv = {part: 5000}
        fcst = {mix: 2000}
        blog = {mix: {"lead_time": 17, "cancel_window": date(2026, 5, 12),
                      "blog_ttl": 6600, "monthly": {6: 0, 8: 6600},
                      "더존코드": "131112", "더존업체명": "(주)에이텍", "part": part}}
        ship = {mix: {"담당자": "신성일", "고객": "(주)에이텍", "고객코드": "131112",
                      "part": part, "y2026": 1200}}
        cols, dash, recs = m5.build_records(inv, fcst, blog, ship)
        self.assertEqual(cols, m5.COLUMNS)
        self.assertEqual(dash, m5.DASHBOARD_COLUMNS)
        self.assertEqual(len(recs), 1)
        row = recs[0]
        self.assertEqual(row["고객코드"], "131112")
        self.assertEqual(row["믹스#"], mix)
        self.assertEqual(row["담당자"], "신성일")
        self.assertEqual(row["품번"], part)
        self.assertEqual(row["Q'ty"], 5000)
        self.assertEqual(row["Lead Time"], 17)
        self.assertEqual(row["Cancel Window"], "2026-05-12")
        self.assertEqual(row["Demand Total"], 2000)
        # Balance = 5000 + 6600 - 2000
        self.assertEqual(row["Balance"], 9600)
        self.assertEqual(row["2026년"], 1200)
        self.assertIsNone(row["2023년"])
        self.assertEqual(row["BLOG TTL"], 6600)
        self.assertEqual(row["8월"], 6600)
        self.assertEqual(row["6월"], 0)

    def test_row_universe_is_union(self):
        m_blog = m5.make_mix(1, "A")
        m_ship = m5.make_mix(2, "B")
        m_fcst = m5.make_mix(3, "C")
        cols, dash, recs = m5.build_records(
            {}, {m_fcst: 9},
            {m_blog: {"lead_time": None, "cancel_window": None, "blog_ttl": 1,
                      "monthly": {}, "더존코드": "1", "더존업체명": None, "part": "A"}},
            {m_ship: {"담당자": None, "고객": None, "고객코드": "2", "part": "B", "y2026": 0}},
        )
        mixes = {r["믹스#"] for r in recs}
        self.assertEqual(mixes, {m_blog, m_ship, m_fcst})

    def test_qty_shared_across_mix_with_same_part(self):
        part = m5._norm_part("P1")
        inv = {part: 800}
        blog = {
            m5.make_mix(10, "P1"): {"lead_time": None, "cancel_window": None, "blog_ttl": 0,
                                    "monthly": {}, "더존코드": "10", "더존업체명": None, "part": part},
            m5.make_mix(20, "P1"): {"lead_time": None, "cancel_window": None, "blog_ttl": 0,
                                    "monthly": {}, "더존코드": "20", "더존업체명": None, "part": part},
        }
        _, _, recs = m5.build_records(inv, {}, blog, {})
        for r in recs:
            self.assertEqual(r["Q'ty"], 800)

    def test_sorted_by_part_then_customer(self):
        blog = {
            m5.make_mix(1, "ZEBRA"): {"lead_time": None, "cancel_window": None, "blog_ttl": 0,
                                      "monthly": {}, "더존코드": "1", "더존업체명": "z", "part": "ZEBRA"},
            m5.make_mix(2, "ALPHA"): {"lead_time": None, "cancel_window": None, "blog_ttl": 0,
                                      "monthly": {}, "더존코드": "2", "더존업체명": "a", "part": "ALPHA"},
        }
        _, _, recs = m5.build_records({}, {}, blog, {})
        self.assertEqual(recs[0]["품번"], "ALPHA")
        self.assertEqual(recs[1]["품번"], "ZEBRA")


if __name__ == "__main__":
    unittest.main()
