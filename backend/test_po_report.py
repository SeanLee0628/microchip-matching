"""발주요청서 보고 양식(표1~표4) 회귀 테스트.

기준은 사용자가 빨간 글씨로 코멘트를 단 `발주요청서_보고_2026-09-08.xlsx` 다.
색·병합·숫자서식·월 범위를 그 파일에 맞춰 놓았으므로, 바뀌면 여기서 걸려야 한다.

main.py(로컬) / main_aws.py(배포) 둘 다 건다 — 사본을 따로 두면 갈라진다.
"""
import io
from datetime import date, datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient

import main
import main_aws
import po_report as pr

BACKENDS = [pytest.param(main, id="main"), pytest.param(main_aws, id="main_aws")]

# 데이터 양식에서 실제로 쓰는 열만 (첨부 파일 4품목 7행과 같은 모양)
COLS = ["CUST PO#", "PO Date", "CRD", "Customer SRD", "FSE", "End customer",
        "DID", "MPN", "BOX_TYPE", "QTY", "DCPL", "SP ($)", "Sales Amt ($)", "PO Customer"]

ROWS = [
    ["MP202608270002", "2026-09-08", "2027-02-04", "2027-02-25", "KATE", "APR",
     "Z42M", "MT53E1G32D2FW-046 WT:B", "TR", 2000, 138.72, 120, 240000, "APR"],
    ["P2602461", "2026-09-08", "2027-01-08", "2027-01-20", "KATE", "RS AUTOMATION",
     "V89C", "MT41K128M16JT-125:K", "TR", 2000, 35.42, 14.6, 29200, "RS AUTOMATION"],
    ["P2602461", "2026-09-08", "2027-02-05", "2027-02-20", "KATE", "RS AUTOMATION",
     "V89C", "MT41K128M16JT-125:K", "TR", 2000, 35.42, 14.6, 29200, "RS AUTOMATION"],
    ["P2602461", "2026-09-08", "2027-01-08", "2027-01-20", "KATE", "RS AUTOMATION",
     "V88A", "MT41K64M16TW-107 AIT:J", "TR", 2000, 19.06, 15.7, 31400, "RS AUTOMATION"],
    ["P2602461", "2026-09-08", "2027-02-05", "2027-02-20", "KATE", "RS AUTOMATION",
     "V88A", "MT41K64M16TW-107 AIT:J", "TR", 2000, 19.06, 15.7, 31400, "RS AUTOMATION"],
    ["P2602461", "2026-09-08", "2027-01-08", "2027-01-20", "KATE", "RS AUTOMATION",
     "QLHS", "MT25QL128ABA8ESF-0SIT", "TR", 1000, 9.92, 8.8, 8800, "RS AUTOMATION"],
    ["P2602461", "2026-09-08", "2027-02-05", "2027-02-20", "KATE", "RS AUTOMATION",
     "QLHS", "MT25QL128ABA8ESF-0SIT", "TR", 1000, 9.92, 8.8, 8800, "RS AUTOMATION"],
]


def _data_bytes(rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(COLS)
    for r in (rows if rows is not None else ROWS):
        ws.append(list(r))
    for n in range(2, ws.max_row + 1):
        for col in (2, 3, 4):  # PO Date / CRD / Customer SRD 를 진짜 날짜로
            c = ws.cell(row=n, column=col)
            if isinstance(c.value, str):
                c.value = datetime.strptime(c.value, "%Y-%m-%d")
            c.number_format = "mm-dd-yy"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _preview(mod, rows=None):
    res = TestClient(mod.app).post(
        "/api/po-report/preview",
        files={"file": ("d.xlsx", _data_bytes(rows), "application/xlsx")})
    assert res.status_code == 200, res.text
    body = res.json()
    assert "error" not in body, body
    return body


def _book(mod, rows=None):
    js = _preview(mod, rows)
    res = TestClient(mod.app).post("/api/po-report/export", json=js)
    assert res.status_code == 200, res.text
    assert res.content[:2] == b"PK"
    return openpyxl.load_workbook(io.BytesIO(res.content))


def _bg(c):
    f = c.fill
    return str(f.fgColor.rgb) if f and f.fgColor and f.fgColor.type == "rgb" else None


def _fg(c):
    """글자색. 지정 안 하면 None (openpyxl 기본 폰트는 color 가 없다)."""
    col = c.font.color if c.font else None
    return str(col.rgb) if col and col.type == "rgb" else None


# ---------- 월 계산 ----------

def test_sales_plan_labels_rolls_over_year():
    """2026-09 기준이면 ~+4M 는 다음 해 1월이다."""
    from datetime import date
    assert pr.sales_plan_labels(date(2026, 9, 8)) == [
        "M(Sep)", "+1M(Oct)", "+2M(Nov)", "+3M(Dec)", "~+4M (Jan~)"]


def test_month_window_is_order_month_to_last_new_po_plus_6():
    rows = _preview(main)["t1_rows"]
    months = pr.month_window(rows)
    assert [pr.month_abbr(m) for m in months] == [
        "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug"]
    assert len(months) == 12, "New PO 마지막(2027-02) + 6개월 = 2027-08"


def test_month_window_shrinks_when_new_po_ends_early():
    """New PO 가 이른 달에 끝나면 창도 그만큼 짧아진다(항상 12개월이 아니다)."""
    rows = [list(ROWS[0])]
    rows[0][2] = "2026-10-05"  # CRD 를 10월로
    months = pr.month_window(_preview(main, rows)["t1_rows"])
    assert [pr.month_abbr(m) for m in months] == [
        "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr"]


# ---------- 표1 ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_t1_customer_srd_is_raw_date_not_korean_month(mod):
    """'2027년 2월' 이 아니라 데이터 AE열 날짜 그대로."""
    rows = _preview(mod)["t1_rows"]
    assert [r["고객사 요청일"] for r in rows[:3]] == ["2027-02-25", "2027-01-20", "2027-02-20"]
    ws = _book(mod)["표1"]
    assert ws["D3"].value == datetime(2027, 2, 25), ws["D3"].value
    assert ws["D3"].number_format == pr.FMT_DATE


@pytest.mark.parametrize("mod", BACKENDS)
def test_t1_header_is_two_rows_with_stacked_crd(mod):
    ws = _book(mod)["표1"]
    assert ws["C1"].value == "Unitrontech" and ws["C2"].value == "CRD"
    assert ws["O1"].value == "Remark"
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "A1:A2" in merged and "O1:O2" in merged
    assert "C1:C2" not in merged, "Unitrontech/CRD 는 2단으로 쌓아야 한다"


@pytest.mark.parametrize("mod", BACKENDS)
def test_t1_header_colors(mod):
    ws = _book(mod)["표1"]
    assert _bg(ws["A1"]) == pr.C_YELLOW
    for coord in ("C1", "K1", "N1"):  # Unitrontech CRD / DCPL / PO Customer
        assert _bg(ws[coord]) == pr.C_GREEN, coord


@pytest.mark.parametrize("mod", BACKENDS)
def test_t1_number_formats_have_comma_and_dollar(mod):
    ws = _book(mod)["표1"]
    assert ws["J3"].number_format == pr.FMT_QTY
    for coord in ("K3", "L3", "M3"):
        assert ws[coord].number_format == pr.FMT_USD, coord


@pytest.mark.parametrize("mod", BACKENDS)
def test_t1_sum_row_keeps_formats(mod):
    ws = _book(mod)["표1"]
    row = len(ROWS) + 3
    assert ws.cell(row=row, column=10).value == 12000
    assert ws.cell(row=row, column=10).number_format == pr.FMT_QTY
    assert ws.cell(row=row, column=13).value == pytest.approx(378800)
    assert ws.cell(row=row, column=13).number_format == pr.FMT_USD


# ---------- 표2 ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_t2_month_abbr_auto_filled_from_date(mod):
    ws = _book(mod)["표2"]
    assert [ws.cell(row=2, column=c).value for c in range(9, 14)] == [
        "M(Sep)", "+1M(Oct)", "+2M(Nov)", "+3M(Dec)", "~+4M (Jan~)"]
    assert ws["I1"].value == "매출 계획"
    assert "I1:M1" in {str(m) for m in ws.merged_cells.ranges}


@pytest.mark.parametrize("mod", BACKENDS)
def test_t2_header_colors_and_sublabels(mod):
    ws = _book(mod)["표2"]
    assert _bg(ws["A1"]) == pr.C_BLUE
    assert _bg(ws["I1"]) == pr.C_LIME
    assert [ws.cell(row=2, column=c).value for c in range(5, 9)] == ["수량", "수량", "Total", "Total"]


@pytest.mark.parametrize("mod", BACKENDS)
def test_t2_blank_cells_carry_comma_format(mod):
    """빈칸에 나중에 숫자를 넣어도 콤마가 붙어야 한다."""
    ws = _book(mod)["표2"]
    assert ws["G3"].value is None
    assert ws["G3"].number_format == pr.FMT_ACCT


# ---------- 표3 ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_blocks_titled_did_mpn(mod):
    ws = _book(mod)["표3"]
    titles = [ws.cell(row=r, column=2).value for r in (1, 10, 19, 28)]
    assert titles == ["(Z42M)MT53E1G32D2FW-046 WT:B", "(V89C)MT41K128M16JT-125:K",
                      "(V88A)MT41K64M16TW-107 AIT:J", "(QLHS)MT25QL128ABA8ESF-0SIT"]


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_new_po_lands_on_crd_month(mod):
    """표1 CRD 월 기준으로 수량이 꽂힌다 — 첨부 파일과 같은 자리·같은 값."""
    ws = _book(mod)["표3"]
    # 블록1(Z42M): CRD 2027-02 → Feb(H) 2000
    assert ws["H7"].value == 2000
    assert ws["G7"].value is None
    # 블록2(V89C): CRD 2027-01, 2027-02 → Jan(G)·Feb(H) 각 2000
    assert (ws["G16"].value, ws["H16"].value) == (2000, 2000)
    # 블록4(QLHS): 1000씩
    assert (ws["G34"].value, ws["H34"].value) == (1000, 1000)


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_manual_rows_left_empty(mod):
    """Delivery/Inventory/Backlog 는 사용자가 채운다 — 자동으로 넣지 않는다."""
    ws = _book(mod)["표3"]
    for r in (4, 5, 6):
        assert all(ws.cell(row=r, column=c).value is None for c in range(3, 15))
        assert ws.cell(row=r, column=3).number_format == pr.FMT_ACCT


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_new_po_row_is_red(mod):
    """자동으로 채워진 New PO 는 수동 입력칸과 구분되게 빨강 (첨부 양식 그대로)."""
    ws = _book(mod)["표3"]
    assert _fg(ws["B7"]) == pr.C_RED
    assert _fg(ws["H7"]) == pr.C_RED
    assert _fg(ws["B4"]) != pr.C_RED, "Delivery 는 검정"


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_balance_is_a_real_formula(mod):
    ws = _book(mod)["표3"]
    assert ws["C8"].value == "=C5-C4+C6+C7", ws["C8"].value
    assert ws["N8"].value == "=N5-N4+N6+N7"
    assert _bg(ws["B8"]) == pr.C_BAL and _bg(ws["C8"]) == pr.C_BAL


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_month_header_row(mod):
    ws = _book(mod)["표3"]
    assert [ws.cell(row=3, column=c).value for c in range(3, 15)] == [
        "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug"]
    assert _bg(ws.cell(row=3, column=3)) == pr.C_MONTH
    assert ws["C1"].value == "소요계획"


# ---------- 2026-09-10 대량발주 요청서 반영분 ----------

def test_fiscal_label_follows_micron_year_starting_september():
    """요청서 표3 FQ 행에서 역산한 규칙: 9월 시작, 연도는 종료연도 표기."""
    f = lambda y, m: pr.fiscal_label(pr._ym(date(y, m, 1)))
    assert f(2026, 3) == "FQ3 YR2026"      # Mar~May
    assert f(2026, 6) == "FQ4 YR2026"      # Jun~Aug
    assert f(2026, 9) == "FQ1 YR2027"      # Sep~Nov, 여기서 회계연도가 넘어간다
    assert f(2026, 12) == "FQ2 YR2027"     # Dec~Feb
    assert f(2027, 1) == "FQ2 YR2027"
    assert f(2027, 3) == "FQ3 YR2027"


def test_fiscal_spans_group_consecutive_months_into_quarters():
    months = [pr._ym(date(2026, 3, 1)) + k for k in range(13)]   # Mar 2026 ~ Mar 2027
    assert pr.fiscal_spans(months) == [
        (0, 3, "FQ3 YR2026"), (3, 3, "FQ4 YR2026"), (6, 3, "FQ1 YR2027"),
        (9, 3, "FQ2 YR2027"), (12, 1, "FQ3 YR2027")]


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_fq_row_sits_above_months_in_every_block(mod):
    """FQ 행은 요청서엔 블록1에만 있었지만 시범이라 보고 전 블록에 넣는다."""
    ws = _book(mod)["표3"]
    for head in (1, 10, 19, 28):
        assert ws.cell(row=head + 1, column=3).value == "FQ1 YR2027", head
        assert ws.cell(row=head + 2, column=3).value == "Sep", head


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_fq_cells_merge_per_quarter(mod):
    ws = _book(mod)["표3"]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"C2:E2", "F2:H2", "I2:K2", "L2:N2"} <= merged, sorted(merged)[:8]
    assert "B1:B3" in merged, "제목은 머리 3행을 덮는다"


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_item_labels_are_centered(mod):
    ws = _book(mod)["표3"]
    assert [ws.cell(row=r, column=2).alignment.horizontal for r in range(4, 9)] == ["center"] * 5


@pytest.mark.parametrize("mod", BACKENDS)
def test_t2_existing_order_is_a_live_formula(mod):
    """기 수주 = Inventory Total + Backlog Total. 값이 아니라 수식으로 고정."""
    ws = _book(mod)["표2"]
    assert ws["E3"].value == "=G3+H3"
    assert ws["E4"].value == "=G4+H4"
    assert ws["E3"].number_format == pr.FMT_ACCT


@pytest.mark.parametrize("mod", BACKENDS)
def test_t2_body_is_malgun_10pt_header_stays_gulim(mod):
    ws = _book(mod)["표2"]
    assert (ws["A3"].font.name, ws["A3"].font.size) == ("맑은 고딕", 10)
    assert ws["A1"].font.name == "굴림", "헤더는 요청서 그대로 굴림"


@pytest.mark.parametrize("mod", BACKENDS)
def test_all_four_tables_have_thin_borders(mod):
    wb = _book(mod)
    for sheet, co in (("표1", "A3"), ("표2", "A3"), ("표3", "B4"), ("표4", "A2")):
        assert wb[sheet][co].border.left.style == "thin", sheet
    assert wb["표1"]["A1"].border.bottom.style == "thin", "헤더도 포함"


@pytest.mark.parametrize("mod", BACKENDS)
def test_t3_gap_row_between_blocks_has_no_border(mod):
    """블록 사이 빈 줄에 선이 가면 표가 하나로 붙어 보인다."""
    ws = _book(mod)["표3"]
    assert ws["B9"].border.left.style is None


# ---------- 표4 (월별 발주 수량) ----------

@pytest.mark.parametrize("mod", BACKENDS)
def test_t4_header_and_rows(mod):
    ws = _book(mod)["표4"]
    assert [ws.cell(row=1, column=j).value for j in (1, 2, 3)] == ["Part (MPN)", "End Customer", "DID"]
    assert ws.cell(row=1, column=4).value == "Jan", "발주 첫 달"
    keys = [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value) for r in range(2, 6)]
    assert keys == [("MT53E1G32D2FW-046 WT:B", "Z42M"),
                    ("MT41K128M16JT-125:K", "V89C"),
                    ("MT41K64M16TW-107 AIT:J", "V88A"),
                    ("MT25QL128ABA8ESF-0SIT", "QLHS")]


def test_t4_window_is_first_po_month_plus_six():
    months, _ = pr.build_t4_rows(_preview(main)["t1_rows"])
    assert [pr.month_abbr(m) for m in months] == [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul"]


def test_t4_window_extends_when_orders_run_past_six_months():
    """뒤를 자르지 않는다 — 발주가 7개월 밖에 있어도 그 달까지 늘린다."""
    rows = [list(ROWS[0]), list(ROWS[1])]
    rows[0][2], rows[1][2] = "2027-01-08", "2027-12-05"
    months, _ = pr.build_t4_rows(_preview(main, rows)["t1_rows"])
    assert [pr.month_abbr(m) for m in months][0] == "Jan"
    assert len(months) == 12 and pr.month_abbr(months[-1]) == "Dec"


def test_t4_sums_orders_split_within_the_same_month():
    """이 표가 필요한 이유 — CRD 를 같은 달에 쪼개 넣은 건을 한 칸으로 합친다."""
    rows = [list(ROWS[1]), list(ROWS[1])]   # 같은 품목·같은 CRD, 2000 씩
    months, t4 = pr.build_t4_rows(_preview(main, rows)["t1_rows"])
    assert len(t4) == 1
    assert t4[0]["qty"][months[0]] == 4000


def test_t4_values_equal_t3_new_po():
    """표4 는 표3 New PO 와 같은 원천 — 수치가 갈리면 안 된다."""
    rows = _preview(main)["t1_rows"]
    _, blocks = pr.build_t3_blocks(rows)
    _, t4 = pr.build_t4_rows(rows)
    by_key = {(r["Part (MPN)"], r["DID"]): r["qty"] for r in t4}
    assert len(by_key) == len(blocks) == 4
    for b in blocks:
        did, mpn = b["title"][1:].split(")", 1)
        assert by_key[(mpn, did)] == b["new_po"], b["title"]


@pytest.mark.parametrize("mod", BACKENDS)
def test_t4_month_without_orders_is_zero_in_accounting_format(mod):
    """0 을 넣어야 회계 서식이 '-' 로 보여준다 (요청서와 같은 모양)."""
    ws = _book(mod)["표4"]
    # 2행 = Z42M, 발주는 CRD 2027-02 한 건뿐 → Feb(E) 에만 2000, 나머지 달은 0
    assert ws["D2"].value == 0, "Jan 은 발주 없음"
    assert ws["E2"].value == 2000
    assert [c.value for c in ws[2][3:]] == [0, 2000, 0, 0, 0, 0, 0]
    assert all(c.number_format == pr.FMT_ACCT for c in ws[2][3:])


@pytest.mark.parametrize("mod", BACKENDS)
def test_export_filename_header_is_latin1_safe(mod):
    js = _preview(mod)
    res = TestClient(mod.app).post("/api/po-report/export", json=js)
    res.headers["content-disposition"].encode("latin-1")
    assert "filename*=UTF-8''" in res.headers["content-disposition"]
