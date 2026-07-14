# -*- coding: utf-8 -*-
"""LOT 시트 파서 테스트 (stdlib unittest + openpyxl).

시트 이름은 파일마다 다르다: Mobis 출고내역 = 'MOBIS shipping lot management',
덕산 납품 LOT# = 'LOT'. 이름이 아니라 구조로 찾아야 한다.
실행: python -m unittest test_labelmaker_lot -v   (backend/ 에서)
"""
import io
import unittest
from datetime import datetime

import openpyxl

import labelmaker_lot as LP


def lot_workbook(sheet_names, lot_sheet, extra_rows=None):
    """sheet_names 순서대로 시트를 만들고, lot_sheet 에만 LOT 블록 구조를 넣는다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheet_names:
        ws = wb.create_sheet(name)
        if name != lot_sheet:
            ws.append(["아무거나", "관계없는", "시트"])
            continue
        # 4열 블록 2개: [입고일, 입고처, DATE CODE, LOT(=MPN 헤더)]
        ws.append(["M3012-0012942", None, None, None, "M3203-001342", None, None, None])
        ws.append(["입고일", "입고처", "DATE CODE", "NEO-M8L-06B",
                   "입고일", "입고처", "DATE CODE", "MT35XU01GBBA1G12"])
        ws.append([datetime(2026, 7, 14), "덕산", "202620", "MN262020446",
                   datetime(2026, 7, 14), "덕산", "202614", "BYHRVVG.41"])
        for row in (extra_rows or []):
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def sheet_workbook(rows, name="LOT"):
    """행 리스트를 그대로 담은 시트 하나짜리 워크북."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestNewFormat(unittest.TestCase):
    """새 포맷: MOBIS ID 행 → 부품번호 행 → 헤더 행(납품일|납품처|DATE CODE|LOT#|Q'ty|MSL) → 데이터.
    부품번호가 LOT 열 헤더가 아니라 별도 행에 있고, 제목 행이 위에 붙기도 한다."""

    NEW = [
        ["변경", None, None, None, None, None],
        ["M3203-001342", None, None, None, None, None],
        ["MT35XU01GBBA1G12-0AAT", None, None, None, None, None],
        ["납품일", "납품처", "DATE CODE", "LOT#", "Q'ty", "MSL"],
        [datetime(2026, 7, 14), "덕산", "202614", "BYHRVVG.41", 1122, 3],
        [datetime(2026, 7, 14), "덕산", "202614", "BY5HQVG.41", 1122, 3],
    ]

    def test_new_format_parses(self):
        lots = LP.parse(sheet_workbook(self.NEW))
        self.assertEqual(len(lots), 2)
        l = lots[0]
        self.assertEqual(l["mobis_id"], "M3203-001342")
        self.assertEqual(l["vpn"], "MT35XU01GBBA1G12-0AAT")   # 헤더가 아니라 윗 행에서
        self.assertEqual(l["lot"], "BYHRVVG.41")              # 헤더가 'LOT#' 인 열
        self.assertEqual(l["datecode"], "202614")
        self.assertEqual(l["qty"], "1122")                    # 헤더가 "Q'ty"
        self.assertEqual(l["msl"], "3")
        self.assertEqual(l["received"], "2026-07-14")

    def test_new_format_without_title_row(self):
        """제목행이 없어도 (MOBIS → 부품번호 → 헤더) 동작."""
        lots = LP.parse(sheet_workbook(self.NEW[1:]))
        self.assertEqual([l["lot"] for l in lots], ["BYHRVVG.41", "BY5HQVG.41"])
        self.assertEqual(lots[0]["vpn"], "MT35XU01GBBA1G12-0AAT")

    def test_old_and_new_side_by_side(self):
        """실제 TEST 파일 모양: 좌측 기존 포맷 + 우측 새 포맷, 헤더 행 높이도 다름."""
        rows = [
            ["기존", None, None, None, None, "변경", None, None, None, None],
            ["M3203-001342", None, None, None, None, "M3203-001342", None, None, None, None],
            ["납품일", "납품처", "DATE CODE", "MT35XU01GBBA1G12-0AAT", None,
             "MT35XU01GBBA1G12-0AAT", None, None, None, None],
            [datetime(2026, 7, 14), "덕산", "202614", "BYHRVVG.41", None,
             "납품일", "납품처", "DATE CODE", "LOT#", "Q'ty"],
            [datetime(2026, 7, 14), "덕산", "202614", "BY5HQVG.41", None,
             datetime(2026, 7, 14), "덕산", "202614", "BYHRVVG.41", 1122],
        ]
        lots = LP.parse(sheet_workbook(rows))
        old = [l for l in lots if l["qty"] == ""]
        new = [l for l in lots if l["qty"]]
        self.assertEqual(len(old), 2)                 # 기존 블록: LOT 2건, 수량 없음
        self.assertEqual(len(new), 1)                 # 새 블록: LOT 1건, 수량 1122
        self.assertEqual(new[0]["qty"], "1122")
        self.assertEqual(new[0]["vpn"], "MT35XU01GBBA1G12-0AAT")
        self.assertEqual({l["vpn"] for l in lots}, {"MT35XU01GBBA1G12-0AAT"})


class TestQtyAndMsl(unittest.TestCase):
    """수량·MSL 열이 파일에 있으면 읽는다. 없으면 빈 값 (화면에서 수기 입력 / MSL 기본 3)."""

    def test_six_column_block_reads_qty_and_msl(self):
        lots = LP.parse(sheet_workbook([
            ["M3012-0012942", None, None, None, None, None],
            ["납품일", "납품처", "DATE CODE", "NEO-M8L-06B", "수량", "MSL"],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020446", 2500, 3],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020442", 1200, 3],
        ]))
        self.assertEqual([l["qty"] for l in lots], ["2500", "1200"])
        self.assertEqual([l["msl"] for l in lots], ["3", "3"])
        self.assertEqual(lots[0]["lot"], "MN262020446")
        self.assertEqual(lots[0]["vpn"], "NEO-M8L-06B")

    def test_four_column_block_leaves_qty_and_msl_empty(self):
        """하위호환: 기존 4열 파일은 수량·MSL 이 빈 값."""
        lots = LP.parse(lot_workbook(["LOT"], "LOT"))
        self.assertEqual([l["qty"] for l in lots], ["", ""])
        self.assertEqual([l["msl"] for l in lots], ["", ""])

    def test_mixed_block_widths(self):
        """자재마다 열 구성이 달라도 된다 (6열 블록 + 4열 블록 혼재)."""
        lots = LP.parse(sheet_workbook([
            ["M3012-0012942", None, None, None, None, None, "M3203-001342", None, None, None],
            ["납품일", "납품처", "DATE CODE", "NEO-M8L-06B", "수량", "MSL",
             "납품일", "납품처", "DATE CODE", "MT35XU01GBBA1G12-0AAT"],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020446", 2500, 3,
             datetime(2026, 7, 14), "덕산", "202614", "BYHRVVG.41"],
        ]))
        by_vpn = {l["vpn"]: l for l in lots}
        self.assertEqual(by_vpn["NEO-M8L-06B"]["qty"], "2500")
        self.assertEqual(by_vpn["NEO-M8L-06B"]["msl"], "3")
        self.assertEqual(by_vpn["MT35XU01GBBA1G12-0AAT"]["qty"], "")     # 그 블록엔 수량 열이 없다
        self.assertEqual(by_vpn["MT35XU01GBBA1G12-0AAT"]["lot"], "BYHRVVG.41")

    def test_header_aliases(self):
        """헤더 이름이 QTY / MSL LEVEL 이어도 인식."""
        lots = LP.parse(sheet_workbook([
            ["M3012-0012942", None, None, None, None, None],
            ["입고일", "입고처", "DATE CODE", "NEO-M8L-06B", "QTY", "MSL LEVEL"],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020446", "1,000", "3"],
        ]))
        self.assertEqual(lots[0]["qty"], "1000")     # 콤마 제거
        self.assertEqual(lots[0]["msl"], "3")

    def test_blank_qty_cell_is_empty_string(self):
        """수량 열은 있는데 그 행만 비었으면 빈 값 (화면에서 수기 입력)."""
        lots = LP.parse(sheet_workbook([
            ["M3012-0012942", None, None, None, None, None],
            ["납품일", "납품처", "DATE CODE", "NEO-M8L-06B", "수량", "MSL"],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020446", None, None],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020442", 800, 2],
        ]))
        self.assertEqual([l["qty"] for l in lots], ["", "800"])
        self.assertEqual([l["msl"] for l in lots], ["", "2"])

    def test_float_qty_becomes_integer_string(self):
        """엑셀이 2500 을 2500.0 으로 주더라도 라벨에는 2500 으로."""
        lots = LP.parse(sheet_workbook([
            ["M3012-0012942", None, None, None, None],
            ["납품일", "납품처", "DATE CODE", "NEO-M8L-06B", "수량"],
            [datetime(2026, 7, 14), "덕산", "202620", "MN262020446", 2500.0],
        ]))
        self.assertEqual(lots[0]["qty"], "2500")


class TestSheetLookup(unittest.TestCase):
    def test_lot_sheet_named_LOT_parses(self):
        """덕산 납품 파일: 시트명이 'LOT' 이어도 파싱된다 (예전엔 ValueError)."""
        lots = LP.parse(lot_workbook(["LOT"], "LOT"))
        self.assertEqual(len(lots), 2)
        self.assertEqual(lots[0]["mobis_id"], "M3012-0012942")
        self.assertEqual(lots[0]["vpn"], "NEO-M8L-06B")
        self.assertEqual(lots[0]["lot"], "MN262020446")
        self.assertEqual(lots[0]["datecode"], "202620")
        self.assertEqual(lots[0]["received"], "2026-07-14")

    def test_original_mobis_sheet_still_parses(self):
        """원본 Mobis 출고내역 시트명도 그대로 동작한다."""
        lots = LP.parse(lot_workbook([LP.SHEET], LP.SHEET))
        self.assertEqual(len(lots), 2)
        self.assertEqual(lots[1]["mobis_id"], "M3203-001342")

    def test_named_sheet_wins_over_structural_match(self):
        """이름이 맞는 시트가 있으면 그걸 쓴다 (다른 시트가 구조상 비슷해도)."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        decoy = wb.create_sheet("LOT")
        decoy.append(["M9999-999999", None, None, None])
        decoy.append(["입고일", "입고처", "DATE CODE", "DECOY-MPN"])
        decoy.append([datetime(2020, 1, 1), "x", "202001", "DECOYLOT"])
        real = wb.create_sheet(LP.SHEET)
        real.append(["M3012-0012942", None, None, None])
        real.append(["입고일", "입고처", "DATE CODE", "NEO-M8L-06B"])
        real.append([datetime(2026, 7, 14), "덕산", "202620", "MN262020446"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        lots = LP.parse(buf)
        self.assertEqual([l["vpn"] for l in lots], ["NEO-M8L-06B"])

    def test_no_lot_sheet_raises_with_sheet_list(self):
        """LOT 구조가 아예 없으면 에러 — 있는 시트 목록을 알려준다."""
        wb = openpyxl.Workbook()
        wb.active.title = "요약"
        wb.active.append(["합계", 1, 2])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        with self.assertRaises(ValueError) as cm:
            LP.parse(buf)
        self.assertIn("요약", str(cm.exception))


class TestLotCell(unittest.TestCase):
    def test_datecode_embedded_in_lot(self):
        self.assertEqual(LP.parse_lot_cell("201442(BYM1JQQ.21)", None),
                         ("201442", "BYM1JQQ.21"))

    def test_plain_lot_uses_datecode_column(self):
        self.assertEqual(LP.parse_lot_cell("BYHRVVG.41", "202614"),
                         ("202614", "BYHRVVG.41"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
