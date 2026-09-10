# -*- coding: utf-8 -*-
"""영업5실 매칭 (Uniquant) 순수 로직.

4개 raw 파일(재고·FCST·백록·출고내역)을 MIX#(더존코드+PART#) 기준으로 조인해
매칭 레코드를 만든다. 무상태·DB/AI 미사용. FastAPI 엔드포인트는 main.py/main_aws.py.
"""
import io
import math
import re
from datetime import date, datetime, timedelta, timezone

KST = timezone(timedelta(hours=9))


def _kst_today():
    """프로덕션 EB 서버가 UTC라도 한국 날짜 기준으로 '오늘'을 잡는다."""
    return datetime.now(KST).date()

from openpyxl import load_workbook

COLUMNS = [
    "고객코드", "믹스#", "담당자", "고객", "품번", "Q'ty", "Lead Time",
    "Cancel Window", "Demand Total", "Balance",
    "2023년", "2024년", "2025년", "2026년", "23~25추이", "25-26(w/BL)", "BLOG TTL",
    "6월", "7월", "8월", "9월", "10월", "11월", "12월", "1월", "2월", "3월",
]
DASHBOARD_COLUMNS = [
    "고객코드", "믹스#", "담당자", "고객", "품번", "Q'ty", "Lead Time",
    "Cancel Window", "Demand Total", "Balance",
    "2023년", "2024년", "2025년", "2026년", "BLOG TTL",
]
MONTH_COLUMNS = [("6월", 6), ("7월", 7), ("8월", 8), ("9월", 9), ("10월", 10),
                 ("11월", 11), ("12월", 12), ("1월", 1), ("2월", 2), ("3월", 3)]
DEFAULT_INVENTORY_PASSWORD = "9178"

# 백록 파일 신선도 경고 임계(일). 기준일은 시트명(YYMMDD)에서 읽는다.
BLOG_STALE_DAYS = 7


def month_window(base_year):
    """월 컬럼(6월…3월)이 실제로 가리키는 (연,월) 목록.

    컬럼 라벨에는 연도가 없어서 예전에는 PDD 의 '월'만 보고 담았다. 그 결과
    2027-07·2028-06 같은 먼 미래 물량이 '7월'·'6월' 칸에 섞여 들어갔다.
    라벨 순서(6→12→1→3)가 곧 시간 순서이므로 월이 되감기는 지점에서 연도를 올린다.
    반환: {(year, month): 라벨}
    """
    out = {}
    year = base_year
    prev = None
    for label, mnum in MONTH_COLUMNS:
        if prev is not None and mnum < prev:
            year += 1
        out[(year, mnum)] = label
        prev = mnum
    return out


def _s(v):
    """str 변환, NaN/None/'nan' 은 None."""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    s = str(v).strip()
    return s if s and s.lower() != "nan" else None


def _code_str(v):
    """더존코드: float(131112.0) → '131112'."""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def _norm_part(p):
    """PART# 정규화: 모든 공백/줄바꿈 제거 + 대문자."""
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return None
    s = re.sub(r"\s+", "", str(p)).upper()
    return s or None


def make_mix(code, part):
    """MIX# = 정규화(더존코드) + 정규화(PART#). 둘 중 하나라도 없으면 None."""
    c = _code_str(code)
    p = _norm_part(part)
    if not c or not p:
        return None
    return c + p


def _to_float(v):
    if v is None:
        return None
    try:
        f = float(v)
        if math.isnan(f) or math.isinf(f):
            return None
        return f
    except (ValueError, TypeError):
        return None


def _to_date(v):
    """datetime/date/Timestamp → date. 8자리 YYYYMMDD(int/str)도 인식. 그 외 None.

    openpyxl이 날짜 시리얼 범위를 벗어났다고 판단해 20250825 같은 8자리 정수/문자열로
    넘기는 셀이 실제로 존재 → 출하 2026 합산·PDD 월버킷에서 조용히 누락되는 것을 방지.
    """
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, str)) and not isinstance(v, bool):
        s = str(v).strip()
        if len(s) == 8 and s.isdigit():
            try:
                return date(int(s[:4]), int(s[4:6]), int(s[6:8]))
            except ValueError:
                return None
    return None


def _find_header_row(ws, required, max_scan=6):
    """required(집합) 의 헤더 텍스트를 모두 포함하는 첫 행을 찾는다.
    반환: (행번호(1-base), {헤더텍스트: 열번호(1-base)}). 못 찾으면 (None, {}).
    헤더 비교는 공백·줄바꿈 제거 후 정확 일치.
    """
    def norm(v):
        if v is None:
            return None
        return re.sub(r"\s+", "", str(v))

    req_norm = {norm(x) for x in required}
    for r in range(1, min(ws.max_row, max_scan) + 1):
        colmap = {}
        present = set()
        for c in range(1, ws.max_column + 1):
            key = norm(ws.cell(row=r, column=c).value)
            if key is None:
                continue
            colmap.setdefault(key, c)
            if key in req_norm:
                present.add(key)
        if req_norm.issubset(present):
            # 원본 헤더 텍스트(required) 기준으로 열번호 매핑 반환
            out = {}
            for orig in required:
                out[orig] = colmap.get(norm(orig))
            return r, out
    return None, {}


def _open_first_ws(contents):
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    return wb[wb.sheetnames[0]]


def _sheet_asof(title):
    """백록 시트명(예 '260609' / '20260609') → date. 그 형식이 아니면 None."""
    s = re.sub(r"\D", "", str(title or ""))
    try:
        if len(s) == 6:
            return date(2000 + int(s[:2]), int(s[2:4]), int(s[4:6]))
        if len(s) == 8:
            return date(int(s[:4]), int(s[4:6]), int(s[6:8]))
    except ValueError:
        return None
    return None


def _pick_base_year(by_ym, fallback_year):
    """월 컬럼 창(6월~3월)의 시작 연도를 고른다.
    후보 연도별로 창 안에 들어오는 물량을 세어 가장 많이 담기는 연도를 쓴다.
    (컬럼 라벨에 연도가 없으므로 데이터로 정렬시키는 것 외에 방법이 없다.)"""
    if not by_ym:
        return fallback_year
    years = {y for (y, _m) in by_ym} | {fallback_year}
    best, best_qty = None, -1.0
    for cand in sorted(years):
        win = month_window(cand)
        covered = sum(q for ym, q in by_ym.items() if ym in win)
        if covered > best_qty:
            best, best_qty = cand, covered
    return best


def read_blog(contents, today=None):
    """백록(Blog) → (records, meta).

    records = {mix: {lead_time, cancel_window(date), blog_ttl, monthly{month:qty},
                     out_of_window, 더존코드, 더존업체명, part}}
    - blog_ttl 은 PDD 와 무관한 Qty Due 전체 합계.
    - monthly 는 월 컬럼 창(6월~3월) **안에 실제로 해당하는 연월**만 담는다.
      창 밖(예 2027-04, 2028-06)은 out_of_window 로 따로 빼고 meta 에 상세를 남긴다.
    - Cancel Window = (PDD - ChangeWindow) 중 today(기본 KST 오늘) 이상 중 가장 가까운 날짜.
    meta = {sheet, asof, base_year, rows, qty_total, by_ym, out_of_window}
    """
    if today is None:
        today = _kst_today()
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb[wb.sheetnames[0]]
    meta = {"sheet": ws.title, "asof": _sheet_asof(ws.title), "base_year": None,
            "rows": 0, "qty_total": 0.0, "by_ym": {}, "out_of_window": {}}

    required = {"PART#", "더존코드", "Qty Due", "PDD", "Change Window",
                "Lead Time Weeks", "Full Manufacturing Cycle Time Weeks", "더존업체명"}
    hdr, col = _find_header_row(ws, required, max_scan=4)
    if hdr is None:
        return {}, meta

    # 1차: 행을 읽어 두고 연월 분포를 만든다(창 시작 연도 결정에 필요).
    parsed = []
    for r in range(hdr + 1, ws.max_row + 1):
        def cv(name):
            c = col.get(name)
            return ws.cell(row=r, column=c).value if c else None

        mix = make_mix(cv("더존코드"), cv("PART#"))
        if not mix:
            continue
        qty = _to_float(cv("Qty Due")) or 0
        pdd = _to_date(cv("PDD"))
        parsed.append({
            "mix": mix, "qty": qty, "pdd": pdd,
            "cw_days": _to_float(cv("Change Window")),
            "lt1": _to_float(cv("Lead Time Weeks")),
            "lt2": _to_float(cv("Full Manufacturing Cycle Time Weeks")),
            "code": _code_str(cv("더존코드")), "name": _s(cv("더존업체명")),
            "part": _norm_part(cv("PART#")),
        })
        meta["rows"] += 1
        meta["qty_total"] += qty
        if pdd is not None:
            ym = (pdd.year, pdd.month)
            meta["by_ym"][ym] = meta["by_ym"].get(ym, 0) + qty

    ref = meta["asof"] or today
    meta["base_year"] = _pick_base_year(meta["by_ym"], ref.year)
    window = month_window(meta["base_year"])

    out = {}
    for p in parsed:
        rec = out.get(p["mix"])
        if rec is None:
            rec = {"lead_time": None, "cancel_window": None, "blog_ttl": 0,
                   "monthly": {}, "out_of_window": 0, "더존코드": p["code"],
                   "더존업체명": p["name"], "part": p["part"]}
            out[p["mix"]] = rec

        rec["blog_ttl"] += p["qty"]
        pdd = p["pdd"]
        if pdd is not None:
            ym = (pdd.year, pdd.month)
            if ym in window:
                rec["monthly"][pdd.month] = rec["monthly"].get(pdd.month, 0) + p["qty"]
            else:
                rec["out_of_window"] += p["qty"]
                meta["out_of_window"][ym] = meta["out_of_window"].get(ym, 0) + p["qty"]

        lt = max([x for x in (p["lt1"], p["lt2"]) if x is not None], default=None)
        if lt is not None:
            rec["lead_time"] = lt if rec["lead_time"] is None else max(rec["lead_time"], lt)

        if pdd is not None and p["cw_days"] is not None:
            cancel = pdd - timedelta(days=int(p["cw_days"]))
            if cancel >= today:  # 이미 지난 날짜는 제외(KST 오늘 이상만)
                if rec["cancel_window"] is None or cancel < rec["cancel_window"]:
                    rec["cancel_window"] = cancel  # 미래 중 가장 가까운 날짜
    return out, meta


def parse_blog(contents, today=None):
    """read_blog 의 records 만 반환하는 하위호환 래퍼."""
    return read_blog(contents, today=today)[0]


def parse_shipment(contents, cutoff_date=None):
    """출고내역 → {mix: {담당자, 고객, 고객코드, part, y2023, y2024, y2025, y2026}}.
    출고일자의 연도별로 출고수량을 버킷팅(2023·2024·2025·2026).
    cutoff_date 지정 시 출고일자 >= cutoff_date 인 행은 2026 합산에서 제외(없으면 2026 전체)."""
    ws = _open_first_ws(contents)
    required = {"고객코드", "고객", "담당자", "품번", "출고일자", "출고수량"}
    hdr, col = _find_header_row(ws, required, max_scan=4)
    if hdr is None:
        return {}

    out = {}
    for r in range(hdr + 1, ws.max_row + 1):
        def cv(name):
            c = col.get(name)
            return ws.cell(row=r, column=c).value if c else None

        mix = make_mix(cv("고객코드"), cv("품번"))
        if not mix:
            continue
        rec = out.get(mix)
        if rec is None:
            rec = {"담당자": _s(cv("담당자")), "고객": _s(cv("고객")),
                   "고객코드": _code_str(cv("고객코드")), "part": _norm_part(cv("품번")),
                   "y2023": 0, "y2024": 0, "y2025": 0, "y2026": 0}
            out[mix] = rec

        d = _to_date(cv("출고일자"))
        qty = _to_float(cv("출고수량")) or 0
        if d is not None:
            if d.year == 2026:
                if cutoff_date is None or d < cutoff_date:
                    rec["y2026"] += qty
            elif d.year in (2023, 2024, 2025):
                rec[f"y{d.year}"] += qty
    return out


def parse_fcst(contents, sheet_name="Sales Revenue"):
    """FCST Sales Revenue → {mix: {demand, code, part, 담당자, 고객}}.
    MIX = make_mix(Customer Code, MPN). demand=Demand Total 합계(중복행 합산),
    나머지 식별자는 해당 mix 첫 행 기준."""
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    required = {"Customer Code", "MPN", "Demand Total", "담당자", "Customer"}
    hdr, col = _find_header_row(ws, required, max_scan=5)
    if hdr is None:
        return {}

    out = {}
    for r in range(hdr + 1, ws.max_row + 1):
        def cv(name):
            c = col.get(name)
            return ws.cell(row=r, column=c).value if c else None

        mix = make_mix(cv("Customer Code"), cv("MPN"))
        if not mix:
            continue
        demand = _to_float(cv("Demand Total"))
        if demand is None:
            continue
        rec = out.get(mix)
        if rec is None:
            rec = {"demand": 0, "code": _code_str(cv("Customer Code")),
                   "part": _norm_part(cv("MPN")), "담당자": _s(cv("담당자")),
                   "고객": _s(cv("Customer"))}
            out[mix] = rec
        rec["demand"] += demand
    return out


def _sum_available_qty(ws):
    """재고 워크시트 → {part_norm: available Q'ty 합계}.  헤더 못 찾으면 {}."""
    hdr, col = _find_header_row(ws, {"Part#", "available Q'ty"}, max_scan=4)
    if hdr is None:
        return {}
    out = {}
    pc, qc = col["Part#"], col["available Q'ty"]
    for r in range(hdr + 1, ws.max_row + 1):
        part = _norm_part(ws.cell(row=r, column=pc).value)
        if not part:
            continue
        qty = _to_float(ws.cell(row=r, column=qc).value) or 0
        out[part] = out.get(part, 0) + qty
    return out


def parse_inventory(contents, password=DEFAULT_INVENTORY_PASSWORD,
                    sheet_name="Jun inventory"):
    """암호화된 재고 .xlsx 를 복호화 후 PART#별 available Q'ty 합계.
    복호화 실패 시 ValueError('재고 파일 비밀번호가 올바르지 않습니다')."""
    import msoffcrypto
    dec = io.BytesIO()
    try:
        of = msoffcrypto.OfficeFile(io.BytesIO(contents))
        of.load_key(password=password)
        of.decrypt(dec)
    except Exception:
        # 암호화가 아닐 수도 있음 → 원본 그대로 시도
        dec = io.BytesIO(contents)
    dec.seek(0)
    try:
        wb = load_workbook(dec, data_only=True)
    except Exception:
        raise ValueError("재고 파일 비밀번호가 올바르지 않습니다")
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    return _sum_available_qty(ws)


def parse_matching_history(contents):
    """마이크로칩(매칭) 파일 → {mix: {y2023, y2024, y2025}}.  과거 연도 출하이력 소스.
    믹스#·2023년·2024년·2025년 헤더가 모두 있는 시트를 (첫 시트만이 아니라)
    워크북 전체에서 찾는다 — 실데이터가 'Lay out' 등 둘째 이후 시트에 있어도 인식.
    빈 셀은 None 유지(원본도 미출하 연도를 공란으로 둠). 헤더 못 찾으면 {}."""
    required = {"믹스#", "2023년", "2024년", "2025년"}
    wb = load_workbook(io.BytesIO(contents), data_only=True)
    ws = None
    col = {}
    hdr = None
    for sn in wb.sheetnames:
        cand = wb[sn]
        h, c = _find_header_row(cand, required, max_scan=8)
        if h is not None:
            ws, hdr, col = cand, h, c
            break
    if ws is None:
        return {}
    mc = col["믹스#"]
    out = {}
    for r in range(hdr + 1, ws.max_row + 1):
        raw = _s(ws.cell(row=r, column=mc).value)
        if not raw:
            continue
        mix = re.sub(r"\s+", "", raw).upper()  # make_mix 출력과 동일 정규화
        rec = {}
        for y in (2023, 2024, 2025):
            c = col.get(f"{y}년")
            rec[f"y{y}"] = _to_float(ws.cell(row=r, column=c).value) if c else None
        out[mix] = rec
    return out


def build_records(inventory, fcst, blog, shipment, history=None, cutoff_date=None):
    """파서 결과 → (COLUMNS, DASHBOARD_COLUMNS, records[dict]).
    행 집합 = 백록 ∪ 출고내역 ∪ FCST 의 MIX. Q'ty는 PART# 기준 조인.
    출하이력: 2023~2025는 마이크로칩(매칭) 파일(history)에서, 2026은 출고내역에서.
    cutoff_date 는 parse_shipment 에서 이미 반영되므로 여기선 미사용(서명 호환용)."""
    records = []
    history = history or {}
    all_mixes = set(blog) | set(shipment) | set(fcst)

    # 2026 출하는 출고내역 파일에 2026 행이 실제로 있을 때만 채운다(없으면 None='-').
    ship_2026_present = any(rec.get("y2026") for rec in shipment.values())

    for mix in all_mixes:
        bl = blog.get(mix, {})
        sh = shipment.get(mix, {})
        fc = fcst.get(mix, {})
        hi = history.get(mix, {})
        demand = fc.get("demand")

        # FCST 단독 + 무수요(None/0) 행은 잡음 → 건너뜀.
        if mix not in shipment and mix not in blog and not demand:
            continue

        part = sh.get("part") or bl.get("part") or fc.get("part")
        qty = inventory.get(part) if part else None
        # 백록에 이 MIX 자체가 없는 것과 '있는데 물량 0'은 다른 사실이다.
        # 전자를 0 으로 찍으면 매칭 실패가 숫자 0 으로 위장된다(엑셀 시절 IFERROR→"0" 과 같은 실패).
        in_blog = mix in blog
        blog_ttl = bl.get("blog_ttl") or 0  # 산식(Balance)용 숫자값

        # 2026 출하 = 출고내역(라이브). 빈 셀/미존재는 None('-').
        y2026 = sh.get("y2026") if ship_2026_present else None

        # Balance = Q'ty + BLOG TTL - Demand. 셋 다 없으면 None.
        if qty is None and blog_ttl == 0 and demand is None:
            balance = None
        else:
            balance = (qty or 0) + blog_ttl - (demand or 0)

        cancel = bl.get("cancel_window")
        rec = {
            "고객코드": sh.get("고객코드") or bl.get("더존코드") or fc.get("code"),
            "믹스#": mix,
            "담당자": sh.get("담당자") or fc.get("담당자"),
            "고객": sh.get("고객") or bl.get("더존업체명") or fc.get("고객"),
            "품번": part,
            "Q'ty": qty,
            "Lead Time": bl.get("lead_time"),
            "Cancel Window": cancel.strftime("%Y-%m-%d") if cancel else None,
            "Demand Total": demand,
            "Balance": balance,
            "2023년": hi.get("y2023"), "2024년": hi.get("y2024"), "2025년": hi.get("y2025"),
            "2026년": y2026,
            "23~25추이": None, "25-26(w/BL)": None,
            "BLOG TTL": blog_ttl if in_blog else None,
        }
        monthly = bl.get("monthly", {})
        for label, mnum in MONTH_COLUMNS:
            rec[label] = monthly.get(mnum, 0) if in_blog else None
        records.append(rec)

    records.sort(key=lambda r: ((r.get("품번") or ""), (r.get("고객") or "")))
    return COLUMNS, DASHBOARD_COLUMNS, records


def diagnostics(meta, records, today=None):
    """검증 경고 목록(문자열). 계산을 막지는 않되, 조용히 지나가던 것을 드러낸다.

    - 신선도: 백록 시트명(YYMMDD)에서 읽은 기준일이 BLOG_STALE_DAYS 초과면 경고
    - 합계 대조: 보고표 BLOG TTL 합 · 월 컬럼 합 + 창 밖 물량 = 백록 원본 Qty Due 총계
    - 월 컬럼 기간 밖 물량 상세(연월별)
    - 백록 미등재로 '-' 처리된 행 수
    """
    if today is None:
        today = _kst_today()
    warn = []
    if not meta:
        return warn

    asof = meta.get("asof")
    if asof is None:
        warn.append(f"백록 기준일을 시트명에서 읽지 못했습니다(시트명 '{meta.get('sheet')}') "
                    f"→ 신선도 검사를 건너뜁니다.")
    else:
        age = (today - asof).days
        if age > BLOG_STALE_DAYS:
            warn.append(f"⚠ 백록 파일이 오래되었습니다: 기준일 {asof.strftime('%Y-%m-%d')} "
                        f"(오늘 기준 {age}일 경과, 임계 {BLOG_STALE_DAYS}일)")

    month_sum = 0.0
    ttl_sum = 0.0
    for r in records:
        for label, _m in MONTH_COLUMNS:
            v = r.get(label)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                month_sum += v
        v = r.get("BLOG TTL")
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            ttl_sum += v

    total = meta.get("qty_total", 0.0)
    out_map = meta.get("out_of_window", {})
    out_qty = sum(out_map.values())

    if round(ttl_sum, 3) != round(total, 3):
        warn.append(f"⚠ 합계 불일치: 보고표 BLOG TTL 합 {ttl_sum:,.0f} ≠ "
                    f"백록 원본 Qty Due 총계 {total:,.0f} (차 {ttl_sum - total:,.0f})")
    if round(month_sum + out_qty, 3) != round(total, 3):
        warn.append(f"⚠ 합계 불일치: 월 컬럼 합 {month_sum:,.0f} + 기간 밖 {out_qty:,.0f} ≠ "
                    f"원본 총계 {total:,.0f} (차 {month_sum + out_qty - total:,.0f})")

    if out_qty:
        base = meta.get("base_year")
        win = f"{base}-06~{base + 1}-03" if base else "-"
        detail = ", ".join(f"{y}-{m:02d} {q:,.0f}"
                           for (y, m), q in sorted(out_map.items()))
        warn.append(f"월 컬럼 기간({win}) 밖 물량 {out_qty:,.0f} 은 월별 칸에 넣지 않았습니다"
                    f"(BLOG TTL 에는 포함): {detail}")

    n_no_blog = sum(1 for r in records if r.get("BLOG TTL") is None)
    if n_no_blog:
        warn.append(f"백록에 없는 MIX {n_no_blog}행 → BLOG TTL·월별 칸을 0 이 아니라 "
                    f"'-'(빈칸)으로 표시했습니다")
    return warn


def export_workbook(columns, rows):
    """columns/rows → 스타일 적용 xlsx(BytesIO).  1행 그룹헤더 + 2행 컬럼헤더 + 데이터."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "마이크로칩(매칭)"

    # 1행: 그룹 헤더
    def col_idx(name):
        return columns.index(name) + 1 if name in columns else None

    y2023 = col_idx("2023년")
    y2026 = col_idx("2026년")
    jun = col_idx("6월")
    mar = col_idx("3월")
    if y2023:
        ws.cell(row=1, column=y2023, value="출하이력")
        if y2026 and y2026 > y2023:
            ws.merge_cells(start_row=1, start_column=y2023, end_row=1, end_column=y2026)
    if jun:
        ws.cell(row=1, column=jun, value="BLOG 2026(PDD기준)")
        if mar and mar > jun:
            ws.merge_cells(start_row=1, start_column=jun, end_row=1, end_column=mar)

    # 2행: 컬럼 헤더
    for i, name in enumerate(columns, start=1):
        ws.cell(row=2, column=i, value=name)

    # 3행~: 데이터 (수치 컬럼은 천단위 쉼표 표시)
    numeric_cols = {
        "Q'ty", "Lead Time", "Demand Total", "Balance",
        "2023년", "2024년", "2025년", "2026년", "23~25추이", "25-26(w/BL)", "BLOG TTL",
        "6월", "7월", "8월", "9월", "10월", "11월", "12월", "1월", "2월", "3월",
    }
    for ridx, row in enumerate(rows, start=3):
        for i, name in enumerate(columns, start=1):
            val = row.get(name)
            cell = ws.cell(row=ridx, column=i, value=val)
            if name in numeric_cols and isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = "#,##0"

    # 스타일: 헤더(1·2행) 하늘색+볼드+가운데
    sky = PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for r in (1, 2):
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = sky
            cell.font = bold
            cell.alignment = center

    # 필터(컬럼 헤더 행부터)
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A2:{last_col}2"
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
